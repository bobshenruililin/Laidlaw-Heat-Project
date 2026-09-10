#!/usr/bin/env python3
"""Build the 10 September Hogan-format live manuscript and supplement.

Wording authorities:
  manuscript/live_collaborative/Heat_CVD_Manuscript_live_update.md
  manuscript/live_collaborative/supplement_live_track.md

The A4/Times/pagination/comment contract follows the archived 24 August
builder. This script does not read governed panels or fit health models.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
LIVE_DIR = ROOT / "manuscript" / "live_collaborative"
MAIN_MD = LIVE_DIR / "Heat_CVD_Manuscript_live_update.md"
SUPP_MD = LIVE_DIR / "supplement_live_track.md"
MAIN_DOCX = LIVE_DIR / "Heat_CVD_Manuscript_20260910_hogan.docx"
MAIN_PDF = LIVE_DIR / "Heat_CVD_Manuscript_20260910_hogan.pdf"
SUPP_DOCX = LIVE_DIR / "Heat_CVD_Supplement_20260910_hogan.docx"
SUPP_PDF = LIVE_DIR / "Heat_CVD_Supplement_20260910_hogan.pdf"

HOGAN_OPEN = "Meteorological data was obtained from the HKO."
HOGAN_AVG = (
    "All monthly data were derived by taking the average of daily data in each calendar month."
)

FORBIDDEN_BODY = (
    "Gate 3",
    "Hogan asked",
    "candidate article",
    "health improved",
    "admissions averted",
    "failing left ventricle",
    "SYNTHETIC_THEORY",
)
FORBIDDEN_COMMENTS = (
    "paste",
    "circulate",
    "shared live",
    "pipeline",
    "gate 3",
    "core panel",
    "rscript",
)

TOKEN = re.compile(
    r"(\*\*.+?\*\*|T~min~|T~max~|NO~2~|SO~2~|O~3~|PM~2\.5~|"
    r"\^[^^]+\^|\\\(.+?\\\)|`[^`]+`|\*[^*]+\*)"
)


def set_run_font(
    run,
    *,
    italic: bool = False,
    bold: bool = False,
    size: float = 12,
    color=None,
    superscript: bool = False,
    subscript: bool = False,
) -> None:
    run.italic = italic
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    run.font.superscript = superscript
    run.font.subscript = subscript


def set_paragraph_format(
    paragraph,
    *,
    first_line: bool = True,
    space_after: float = 8,
    space_before: float = 0,
    align: str = "left",
    line_spacing: float | None = None,
) -> None:
    pf = paragraph.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if line_spacing is None:
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    else:
        pf.line_spacing = line_spacing
    pf.first_line_indent = Cm(1.27) if first_line else Cm(0)
    if align == "center":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    finish = OxmlElement("w:fldChar")
    finish.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, finish))
    set_run_font(run, size=10, italic=True)


def configure_document(doc: Document, header_text: str, *, body_size: float) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(body_size)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.clear()
    set_paragraph_format(hp, first_line=False, space_after=0)
    run = hp.add_run(header_text)
    set_run_font(run, italic=True, size=10, color=RGBColor(0x55, 0x55, 0x55))
    tab_run = hp.add_run("\t")
    set_run_font(tab_run, size=10)
    add_page_number(hp)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), "9360")
    tabs.append(tab)
    hp._p.get_or_add_pPr().append(tabs)


def _inline_math(text: str) -> str:
    text = re.sub(r"\\mathrm\{([^}]+)\}", r"\1", text)
    text = text.replace("\\,", "").replace("\\", "")
    return text


def add_rich_runs(paragraph, text: str, *, size: float, italic: bool = False):
    subs = {
        "T~min~": ("T", "min"),
        "T~max~": ("T", "max"),
        "NO~2~": ("NO", "2"),
        "SO~2~": ("SO", "2"),
        "O~3~": ("O", "3"),
        "PM~2.5~": ("PM", "2.5"),
    }
    runs = []
    pos = 0
    for match in TOKEN.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, size=size, italic=italic)
            runs.append(run)
        token = match.group(0)
        if token in subs:
            base, sub = subs[token]
            base_run = paragraph.add_run(base)
            set_run_font(base_run, size=size, italic=italic)
            sub_run = paragraph.add_run(sub)
            set_run_font(sub_run, size=size, italic=italic, subscript=True)
            runs.extend((base_run, sub_run))
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, italic=italic)
            runs.append(run)
        elif token.startswith("^") and token.endswith("^"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, superscript=True, italic=italic)
            runs.append(run)
        elif token.startswith(r"\(") and token.endswith(r"\)"):
            run = paragraph.add_run(_inline_math(token[2:-2]))
            set_run_font(run, size=size, italic=True)
            runs.append(run)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True)
            runs.append(run)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True)
            runs.append(run)
        else:
            run = paragraph.add_run(token)
            set_run_font(run, size=size, italic=italic)
            runs.append(run)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, italic=italic)
        runs.append(run)
    return runs


def add_body(
    doc: Document,
    text: str,
    *,
    size: float,
    first_line: bool = True,
    italic: bool = False,
    align: str = "justify",
    keep_with_next: bool = False,
):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=first_line, space_after=8, align=align)
    p.paragraph_format.keep_with_next = keep_with_next
    runs = add_rich_runs(p, text, size=size, italic=italic)
    return p, runs


def add_heading(
    doc: Document,
    text: str,
    level: int,
    *,
    size: float,
    new_page: bool = False,
):
    p = doc.add_paragraph()
    set_paragraph_format(
        p,
        first_line=False,
        space_before=0 if new_page else (12 if level == 1 else 8),
        space_after=6,
    )
    p.paragraph_format.keep_with_next = True
    if new_page:
        p.paragraph_format.page_break_before = True
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, bold=True, size=size + 2)
    elif level == 2:
        set_run_font(run, bold=True, size=size)
    else:
        set_run_font(run, bold=True, italic=True, size=size)
    return p, run


def add_caption(
    doc: Document,
    text: str,
    *,
    size: float,
    new_page: bool,
):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, space_after=8)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    if new_page:
        p.paragraph_format.page_break_before = True
    runs = add_rich_runs(p, text, size=size, italic=True)
    return p, runs


def _set_cell_margins(cell, top=40, bottom=40, left=55, right=55) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def _row_cant_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _row_repeat_header(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def clean_cell(text: str) -> str:
    text = text.replace("**", "").replace("`", "")
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = text.replace("T~min~", "Tmin").replace("T~max~", "Tmax")
    text = text.replace("NO~2~", "NO₂").replace("SO~2~", "SO₂")
    text = text.replace("O~3~", "O₃").replace("PM~2.5~", "PM₂.₅")
    return text.strip()


def parse_table(lines: list[str]) -> tuple[list[str], list[list[str]]]:
    content = [line for line in lines if set(line.strip()) - set("|:- ")]
    rows = [
        [clean_cell(cell) for cell in line.strip().strip("|").split("|")]
        for line in content
    ]
    if len(rows) < 2:
        raise ValueError(f"invalid markdown table: {lines!r}")
    return rows[0], rows[1:]


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    *,
    font_size: float,
    widths: list[float] | None = None,
) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for child in list(tbl_pr):
        if child.tag == qn("w:tblW"):
            tbl_pr.remove(child)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), "5000")
    tbl_w.set(qn("w:type"), "pct")
    tbl_pr.append(tbl_w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    if widths and len(widths) == len(headers):
        grid = table._tbl.find(qn("w:tblGrid"))
        if grid is not None:
            table._tbl.remove(grid)
        grid = OxmlElement("w:tblGrid")
        for width in widths:
            col = OxmlElement("w:gridCol")
            col.set(qn("w:w"), str(int(width * 1440)))
            grid.append(col)
        table._tbl.tblPr.addnext(grid)

    all_rows = [headers] + rows
    for row_idx, (row, values) in enumerate(zip(table.rows, all_rows)):
        _row_cant_split(row)
        if row_idx == 0:
            _row_repeat_header(row)
        for col_idx, value in enumerate(values):
            cell = row.cells[col_idx]
            cell.text = ""
            _set_cell_margins(cell)
            if widths and len(widths) == len(headers):
                cell.width = Inches(widths[col_idx])
            p = cell.paragraphs[0]
            set_paragraph_format(p, first_line=False, space_after=1, space_before=1)
            p.paragraph_format.keep_together = True
            if row_idx < len(all_rows) - 1:
                p.paragraph_format.keep_with_next = True
            run = p.add_run(value)
            set_run_font(run, bold=row_idx == 0, size=font_size)


def add_picture(doc: Document, path: Path, *, width: float, max_height: float) -> None:
    from PIL import Image

    with Image.open(path) as image:
        px_w, px_h = image.size
    final_width = width
    if final_width * px_h / px_w > max_height:
        final_width = max_height * px_w / px_h
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, space_after=4, align="center")
    p.paragraph_format.keep_together = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(final_width))


def parse_blocks(markdown: str) -> list[tuple[str, list[str]]]:
    lines = markdown.splitlines()
    blocks: list[tuple[str, list[str]]] = []
    idx = 0
    while idx < len(lines):
        line = lines[idx]
        if not line.strip():
            idx += 1
            continue
        if line.startswith("#### "):
            blocks.append(("h3", [line[5:].strip()]))
            idx += 1
            continue
        if line.startswith("### "):
            blocks.append(("h2", [line[4:].strip()]))
            idx += 1
            continue
        if line.startswith("## "):
            blocks.append(("h1", [line[3:].strip()]))
            idx += 1
            continue
        if line.startswith("# "):
            blocks.append(("h0", [line[2:].strip()]))
            idx += 1
            continue
        if line.startswith("!["):
            blocks.append(("image", [line]))
            idx += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while idx < len(lines) and lines[idx].startswith("|"):
                table_lines.append(lines[idx])
                idx += 1
            blocks.append(("table", table_lines))
            continue
        if line.startswith("\\["):
            equation_lines = [line]
            idx += 1
            while idx < len(lines) and not lines[idx].strip().endswith("\\]"):
                equation_lines.append(lines[idx])
                idx += 1
            if idx < len(lines):
                equation_lines.append(lines[idx])
                idx += 1
            blocks.append(("equation", equation_lines))
            continue
        paragraph = [line]
        idx += 1
        while (
            idx < len(lines)
            and lines[idx].strip()
            and not lines[idx].startswith(("#", "|", "![", "\\["))
        ):
            paragraph.append(lines[idx])
            idx += 1
        blocks.append(("paragraph", paragraph))
    return blocks


def _comment(doc: Document, run, text: str) -> None:
    doc.add_comment(run, text, author="Bob Shen", initials="BS")


def _find_run(runs, needle: str):
    return next((run for run in runs if needle in run.text), runs[0] if runs else None)


def _table_widths(caption_text: str, columns: int) -> list[float] | None:
    if caption_text.startswith("Table 1."):
        return [1.25, 0.75, 0.45, 0.80, 0.70, 2.15]
    if caption_text.startswith("Table 2."):
        return [0.90, 1.40, 1.90, 1.90]
    if caption_text.startswith("Table 3."):
        return [1.70, 1.10, 1.10, 1.10, 1.10]
    if columns:
        return [6.1 / columns] * columns
    return None


def _equation_text(_lines: list[str]) -> str:
    return (
        "log E(Yₜ) = log(dₜ) + α + βXₜ + "
        "Σₘ₌₂¹² γₘ I(monthₜ = m) + s(t; 4 df).    (1)"
    )


def build_main() -> None:
    markdown = MAIN_MD.read_text(encoding="utf-8")
    if HOGAN_OPEN not in markdown or HOGAN_AVG not in markdown:
        raise AssertionError("Hogan weather paragraph drifted")
    for phrase in FORBIDDEN_BODY:
        if phrase.lower() in markdown.lower():
            raise AssertionError(f"forbidden main-body phrase: {phrase}")

    doc = Document()
    configure_document(
        doc,
        "Window sensitivity of first CHD/HF hospitalisation",
        body_size=12,
    )
    at_page_start = True
    pending_caption = ""
    section_breaks = {
        "Introduction",
        "Results",
        "Discussion",
        "Strengths and limitations",
        "References",
    }
    references = False

    for kind, lines in parse_blocks(markdown):
        if kind == "h0":
            p = doc.add_paragraph()
            set_paragraph_format(p, first_line=False, space_after=6, align="center")
            run = p.add_run(lines[0])
            set_run_font(run, bold=True, size=16)
            at_page_start = False
            continue

        if kind == "h1":
            text = lines[0]
            references = text == "References"
            new_page = text in section_breaks and not at_page_start
            add_heading(doc, text, 1, size=12, new_page=new_page)
            at_page_start = False
            continue

        if kind in ("h2", "h3"):
            add_heading(doc, lines[0], 2 if kind == "h2" else 3, size=12)
            at_page_start = False
            continue

        if kind == "equation":
            p = doc.add_paragraph()
            set_paragraph_format(p, first_line=False, space_after=8, align="center")
            run = p.add_run(_equation_text(lines))
            set_run_font(run, italic=True, size=12)
            at_page_start = False
            continue

        if kind == "table":
            headers, rows = parse_table(lines)
            add_table(
                doc,
                headers,
                rows,
                font_size=9,
                widths=_table_widths(pending_caption, len(headers)),
            )
            doc.add_page_break()
            at_page_start = True
            pending_caption = ""
            continue

        if kind == "image":
            match = re.search(r"!\[([^\]]*)\]\(([^)]+)\)", lines[0])
            if not match:
                raise ValueError(f"invalid image markdown: {lines[0]}")
            path = (MAIN_MD.parent / match.group(2)).resolve()
            if not path.is_file():
                raise FileNotFoundError(path)
            max_height = 6.0 if "Figure 3" in match.group(1) else 7.4
            add_picture(doc, path, width=6.2, max_height=max_height)
            doc.add_page_break()
            at_page_start = True
            pending_caption = ""
            continue

        text = " ".join(line.strip() for line in lines)
        plain = clean_cell(text)

        if plain.startswith("Running title:"):
            p, runs = add_body(
                doc,
                plain,
                size=11,
                first_line=False,
                italic=True,
                align="center",
            )
            at_page_start = False
            continue

        if plain.startswith("Bob Ruililin Shen"):
            p, runs = add_body(
                doc,
                text,
                size=12,
                first_line=False,
                align="center",
            )
            author_run = _find_run(runs, "Author 2")
            if author_run is not None:
                _comment(
                    doc,
                    author_run,
                    "Author order after the first and last positions remains for "
                    "Professor Bishai and the team to confirm.",
                )
            at_page_start = False
            continue

        if plain.startswith("1 School of Public Health"):
            add_body(
                doc,
                text,
                size=10,
                first_line=False,
                italic=True,
                align="center",
            )
            at_page_start = False
            continue

        if plain.startswith(("Table ", "Figure ")):
            if not at_page_start:
                new_page = True
            else:
                new_page = False
            p, runs = add_caption(doc, plain, size=11, new_page=new_page)
            pending_caption = plain
            if plain.startswith("Table 2.") and runs:
                _comment(
                    doc,
                    runs[0],
                    "The pre-2020 window is contained in the full window. These "
                    "are sensitivity fits, not independent period estimates.",
                )
            at_page_start = False
            continue

        if references and re.match(r"^\d+\.", plain):
            p = doc.add_paragraph()
            set_paragraph_format(p, first_line=False, space_after=3, align="left")
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.75)
            add_rich_runs(p, text, size=11)
            at_page_start = False
            continue

        first_line = not text.startswith("**")
        p, runs = add_body(doc, text, size=12, first_line=first_line)
        low = plain.lower()
        if plain.startswith("Methods.") and runs:
            _comment(
                doc,
                runs[0],
                "The pre-2020 fit is a nested-window sensitivity. No "
                "2020–2023-only Model 1 or exposure-by-period interaction is "
                "reported.",
            )
        if "uw xx-xxx" in low and runs:
            _comment(
                doc,
                _find_run(runs, "UW XX-XXX"),
                "Professor Bishai should replace the IRB placeholder UW XX-XXX "
                "with the approved reference before submission.",
            )
        if plain.startswith("The dependent variable is") and runs:
            _comment(
                doc,
                runs[0],
                "Roro should confirm the written ICD inclusion list, first-event "
                "timing, and inpatient-versus-DAE construction before submission.",
            )
        if HOGAN_OPEN in plain and runs:
            _comment(
                doc,
                runs[0],
                "Hogan's meteorological wording is retained verbatim, including "
                "the averaging sentence.",
            )
        if plain.startswith("Model 2 additionally") and runs:
            _comment(
                doc,
                runs[0],
                "Model 2 and Model 3 are specified; no new coefficients are "
                "added because the governed panels are not present here.",
            )
        if "laboratory measurements were not in the transfer" in low and runs:
            _comment(
                doc,
                runs[0],
                "Laboratory, infection, vaccination, and serology fields require "
                "a new governed extract if Professor Bishai opens that scope.",
            )
        if plain.startswith("The HF cold-day point estimate") and runs:
            _comment(
                doc,
                runs[0],
                "The two windows are nested. Interval overlap is not a formal "
                "test of coefficient difference.",
            )
        at_page_start = False

    # Avoid an empty terminal page created after the last float if the layout
    # changes in future; current manuscript has prose after Figure 3.
    MAIN_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(MAIN_DOCX)
    print(f"wrote {MAIN_DOCX} ({MAIN_DOCX.stat().st_size} bytes)")


def _supp_image_path(raw: str) -> Path:
    return (SUPP_MD.parent / raw).resolve()


def build_supplement() -> None:
    markdown = SUPP_MD.read_text(encoding="utf-8")
    doc = Document()
    configure_document(
        doc,
        "Supplementary information | Window sensitivity of first CHD/HF hospitalisation",
        body_size=10.5,
    )
    at_page_start = True

    for kind, lines in parse_blocks(markdown):
        if kind == "h0":
            p = doc.add_paragraph()
            set_paragraph_format(p, first_line=False, space_after=12, align="center")
            run = p.add_run(lines[0])
            set_run_font(run, bold=True, size=16)
            at_page_start = False
            continue
        if kind in ("h1", "h2", "h3"):
            level = {"h1": 1, "h2": 2, "h3": 3}[kind]
            new_page = kind == "h1" and not at_page_start
            add_heading(doc, lines[0], level, size=10.5, new_page=new_page)
            at_page_start = False
            continue
        if kind == "table":
            headers, rows = parse_table(lines)
            add_table(
                doc,
                headers,
                rows,
                font_size=7.5 if len(headers) >= 5 else 8.5,
                widths=[6.1 / len(headers)] * len(headers),
            )
            at_page_start = False
            continue
        if kind == "image":
            match = re.search(r"!\[([^\]]*)\]\(([^)]+)\)", lines[0])
            if not match:
                raise ValueError(lines[0])
            path = _supp_image_path(match.group(2))
            if not path.is_file():
                raise FileNotFoundError(path)
            add_picture(doc, path, width=6.1, max_height=7.2)
            at_page_start = False
            continue
        if kind == "equation":
            p = doc.add_paragraph()
            set_paragraph_format(p, first_line=False, space_after=6, align="center")
            run = p.add_run(_equation_text(lines))
            set_run_font(run, italic=True, size=10.5)
            at_page_start = False
            continue
        text = " ".join(line.strip() for line in lines)
        add_body(
            doc,
            text,
            size=10.5,
            first_line=not text.startswith("**"),
        )
        at_page_start = False

    SUPP_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(SUPP_DOCX)
    print(f"wrote {SUPP_DOCX} ({SUPP_DOCX.stat().st_size} bytes)")


def _office_binary() -> str:
    binary = shutil.which("soffice") or shutil.which("libreoffice")
    if binary is None:
        raise SystemExit(
            "LibreOffice is required to build the dated PDF. "
            "Install libreoffice-writer and rerun."
        )
    return binary


def convert_pdf(docx_path: Path, pdf_path: Path) -> None:
    binary = _office_binary()
    if pdf_path.exists():
        pdf_path.unlink()
    with tempfile.TemporaryDirectory(prefix=f"lo-{docx_path.stem}-") as profile:
        cmd = [
            binary,
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(pdf_path.parent),
            str(docx_path),
        ]
        proc = subprocess.run(cmd, check=False, capture_output=True, text=True)
    if proc.returncode != 0 or not pdf_path.is_file():
        raise SystemExit(
            f"LibreOffice conversion failed ({proc.returncode}): "
            f"{proc.stdout}\n{proc.stderr}"
        )
    print(f"wrote {pdf_path} ({pdf_path.stat().st_size} bytes)")


def _norm(text: str) -> str:
    return " ".join(text.replace("\u00ad", "").split()).replace("–", "-").replace("—", "-")


def _pdf_pages(path: Path) -> list[str]:
    import pymupdf

    pdf = pymupdf.open(path)
    return [_norm(page.get_text()) for page in pdf]


def review_main() -> None:
    doc = Document(str(MAIN_DOCX))
    section = doc.sections[0]
    assert round(section.page_width.mm, 1) == 210.0
    assert round(section.page_height.mm, 1) == 297.0
    assert all(
        round(value.cm, 2) == 2.54
        for value in (
            section.top_margin,
            section.bottom_margin,
            section.left_margin,
            section.right_margin,
        )
    )
    comments = list(doc.comments)
    assert len(comments) >= 7
    comment_text = " ".join(comment.text for comment in comments).lower()
    for banned in FORBIDDEN_COMMENTS:
        assert banned not in comment_text, banned

    pages = _pdf_pages(MAIN_PDF)
    full = "\n".join(pages)
    for required in (
        "Analysis-window sensitivity",
        HOGAN_OPEN,
        HOGAN_AVG,
        "Table 1. Outcome summary",
        "Table 2. Official-day count ratios",
        "Table 3. Uncertainty ladder",
        "Figure 1. First-event",
        "Figure 2. Official cold days",
        "Figure 3. Model 1 count ratios",
        "1.022 (1.002-1.042)",
        "1.073 (1.006-1.144)",
        "interval overlap is not a test of a window difference",
        "do not show improved cardiovascular health",
        "Supplementary Table S9",
        "UW XX-XXX",
        "None.",
    ):
        assert required in full, required
    for forbidden in FORBIDDEN_BODY:
        assert forbidden.lower() not in full.lower(), forbidden

    def page_with(needle: str) -> int:
        hits = [idx for idx, text in enumerate(pages) if needle in text]
        assert hits, needle
        return hits[0]

    float_pages = [
        page_with("Table 1. Outcome summary"),
        page_with("Table 2. Official-day count ratios"),
        page_with("Table 3. Uncertainty ladder"),
        page_with("Figure 1. First-event"),
        page_with("Figure 2. Official cold days"),
        page_with("Figure 3. Model 1 count ratios"),
    ]
    assert len(set(float_pages)) == len(float_pages), float_pages
    assert "Coronary heart disease" in pages[float_pages[0]]
    assert "Heart failure" in pages[float_pages[0]]
    assert "CHD" in pages[float_pages[1]] and "HF" in pages[float_pages[1]]
    assert "HF mean minimum temperature" in pages[float_pages[2]]
    assert len(pages) >= 18
    print(f"main PDF checks passed ({len(pages)} A4 pages)")


def review_supplement() -> None:
    pages = _pdf_pages(SUPP_PDF)
    full = "\n".join(pages)
    for required in (
        "Supplementary Table S1",
        "Supplementary Table S7",
        "Supplementary Table S8",
        "Supplementary Table S9",
        "Supplementary Table S10",
        "Supplementary Note S1",
        "SYNTHETIC_CALIBRATION",
        "not a health finding",
        "Later years were hotter",
    ):
        assert required in full, required
    assert "health improved" not in full.lower()
    assert len(pages) >= 10
    print(f"supplement PDF checks passed ({len(pages)} A4 pages)")


def main() -> None:
    build_main()
    build_supplement()
    convert_pdf(MAIN_DOCX, MAIN_PDF)
    convert_pdf(SUPP_DOCX, SUPP_PDF)
    review_main()
    review_supplement()


if __name__ == "__main__":
    main()
