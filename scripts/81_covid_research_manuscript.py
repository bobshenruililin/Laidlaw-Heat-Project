#!/usr/bin/env python3
"""Build/audit the COVID exploratory manuscript using only public aggregate outputs.

Run with the bundled document Python. --audit reads no governed health file and
checks scientific tables independently against approved CSVs. PDF rendering uses
the packaged document renderer after this Word build, not a system office app.
The previous builder remains available for historical versions.
"""
from pathlib import Path
import argparse, csv, hashlib, importlib.util, json, re, sys
ROOT=Path(__file__).resolve().parents[1]
M=ROOT/'manuscript/covid_period'
EXPS=['mean_temp','mean_tmax','mean_tmin','hot_nights','very_hot_days','cold_days']
NAMES=['Mean temperature / 1°C','Mean maximum / 1°C','Mean minimum / 1°C','Hot nights / 5 days','Very hot days / 5 days','Cold days / 5 days']
SCENARIOS=['trend_ns3','trend_ns6','trend_ns8','year_fixed_effects','drop_first_12_months','drop_first_24_months','covid_phase_adjusted']
SOURCES=['outputs/tables/cvd_descriptive_annual_totals.csv','outputs/tables/cvd_descriptive_covid_era_means.csv','outputs/tables/cvd_core_robust_estimates.csv','outputs/tables/cvd_trend_depletion_sensitivity.csv','outputs/tables/cvd_core_model_fit.csv']
def read(path):return list(csv.DictReader((ROOT/path).open()))
def ci(row):return f"{float(row['rr']):.3f} ({float(row['rr_low']):.3f}–{float(row['rr_high']):.3f})"
def blocks(text):
 spec=importlib.util.spec_from_file_location('covid_layout',ROOT/'scripts/78_covid_period_manuscript_docx.py'); module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module)
 return module,module.parse_blocks(text)
def audit():
 main=(M/'Manuscript_covid_period_draft.md').read_text();supp=(M/'Supplement_covid_period.md').read_text()
 b,mb=blocks(main);_,sb=blocks(supp);mt=[b.parse_table(x)[1]for k,x in mb if k=='table'];st=[b.parse_table(x)[1]for k,x in sb if k=='table']
 checks=[]
 def ck(name,condition):
  checks.append({'check':name,'passed':bool(condition)})
  if not condition:raise AssertionError(name)
 annual=read(SOURCES[0]);a={(r['outcome'],int(r['year'])):int(r['total_events'])for r in annual}
 ck('supplement annual table matches all 22 approved outcome totals',st[0]==[[str(y),f"{a['chd',y]:,}",f"{a['hf',y]:,}"]for y in range(2013,2024)])
 sens=read(SOURCES[3]);s={(r['outcome'],r['exposure'],r['scenario']):r for r in sens}
 ck('nested and full official panel matches 12 source estimates and intervals',st[3]==[[o.upper(),dict(zip(EXPS,NAMES))[e],ci(s[o,e,'pre_covid']),ci(s[o,e,'baseline_ns4'])]for o in ['chd','hf']for e in EXPS[3:]])
 era=read(SOURCES[1]);periods=[('pre_covid_2013_2019','2013–2019','84'),('covid_2020_2022','2020–2022','36'),('post_reopening_2023','2023','12')]
 ck('disjoint era means and month counts match approved summaries',st[1]==[[label,n]+[f"{float(next(r['mean_events']for r in era if r['outcome']==o and r['period']==key)):.1f}"for o in ['chd','hf']]for key,label,n in periods])
 core=[r for r in read(SOURCES[2])if r['family']=='negative_binomial' and r['offset_policy']=='days_only' and r['model_structure']=='single_exposure'];d={(r['outcome'],r['exposure']):r for r in core if r['se_method']=='NeweyWest_lag6'}
 ck('complete twelve-fit panel and q family match source',st[2]==[[o.upper(),name,ci(d[o,e]),f"{float(d[o,e]['p_value']):.3f}",f"{float(d[o,e]['q_value_core_bh']):.3f}"]for o in ['chd','hf']for e,name in zip(EXPS,NAMES)])
 for i,o in enumerate(['chd','hf']):
  ck(f'{o} all 21 alternative-calendar estimates match source',[row[1:]for row in st[4+i]]==[[ci(s[o,e,key])for e in EXPS[3:]]for key in SCENARIOS])
  expected=[]
  for e,name in zip(EXPS,NAMES):
   row=[name]
   for method in ['model','HC1','NeweyWest_lag3','NeweyWest_lag6']:
    r=next(r for r in core if (r['outcome'],r['exposure'],r['se_method'])==(o,e,method));row.append(f"{float(r['rr_low']):.4f}–{float(r['rr_high']):.4f}")
   expected.append(row)
  ck(f'{o} all uncertainty bounds match source',st[6+i]==expected)
 for o in ['chd','hf']:
  total=sum(a[o,y]for y in range(2013,2024));ck(f'{o} cumulative total in body',f'{total:,}'in main)
  for y in [2020,2023]:
   pct=100*(1-a[o,y]/a[o,2019]);ck(f'{o} {y} percentage versus 2019',f'{pct:.1f}%'in main)
 ck('multiplicity minimum correctly displayed',round(min(float(r['q_value_core_bh'])for r in d.values()),3)==.192 and '0.192'in main)
 ck('near-null HF minimum interval and p value use unrounded source',f"{float(d['hf','mean_tmin']['rr_high']):.10f}"in supp and f"{float(d['hf','mean_tmin']['p_value']):.4f}"in supp)
 for scenario in ['trend_ns8','drop_first_24_months']:
  ck(f'hf cold-day {scenario} diagnostic example matches frozen estimate and interval',ci(s['hf','cold_days',scenario]) in main)
 ck('nominal pointwise intervals and two-sided tests explicitly distinguished from BH adjustment','Nominal pointwise 95%' in main and 'two-sided lag-six Wald p-values for zero exposure coefficient' in main and 'It did not adjust interval endpoints.' in main)
 ck('window lengths not misrepresented as independently verified fitted observations','rather than verified fitted-observation counts' in main and 'rather than the number of observations used by each fitted model' in supp)
 ck('reference numbering complete',sorted(set(int(x)for group in re.findall(r'\[([0-9, ]+)\]',main) for x in group.split(',')))==list(range(1,18)) and len(re.findall(r'^\d+\. ',main,re.M))==17)
 seen=[]
 for group in re.findall(r'\[([0-9, ]+)\]',main.split('## References')[0]):
  for value in group.split(','):
   if int(value)not in seen:seen.append(int(value))
 ck('references follow first-citation order',seen==list(range(1,18)))
 weather=next(x for x in main.split('\n\n')if x.startswith(b.HOGAN_OPEN));live=(ROOT/'manuscript/live_collaborative/Heat_CVD_Manuscript_live_update.md').read_text();ck('entire protected weather paragraph unchanged',weather in live)
 ck('all health source rows retain approved aggregate label',all(r['data_status']=='HA_APPROVED_AGGREGATE'for r in annual+era+core+sens))
 fm=json.loads((ROOT/'figures/covid_period/frozen_publication/publication_figures_manifest.json').read_text())
 ck('all three publication figure source and export hashes match current artifacts',len(fm['figures'])==3 and all(hashlib.sha256((ROOT/item['path']).read_bytes()).hexdigest()==item['sha256'] for f in fm['figures'] for item in f['inputs']+[f['source_snapshot']]+f['outputs']))
 ck('conceptual diagram remains supplementary and explicitly hypothetical','Conceptual diagram' in supp and 'conceptual observation process' not in main.lower())
 ck('three new close precedents are cited without inconsistent abstract numbers',all(doi in main for doi in ['10.1002/hsr2.71291','10.1038/s41440-025-02230-y','10.1111/dom.70984']) and '1,226,685' not in main and '1.108' not in main)
 ck('actual annual counts caption matches the revised figure', 'zero-origin count axes' in main and 'indexed to 2013' not in main)
 ck('new clinical follow-up is external evidence, with less conclusive endpoints shown','results for stroke and all-cause mortality were less conclusive' in main)

 ck('no temporary tokens or fabricated declarations',not any(x in main+supp for x in ['ANNUAL_TABLE','OFFICIAL_TABLE','REFERENCES','UW XX-XXX','Author 2','Gate 3','ruling-out exhibit','Positive serial correlation should inflate']))
 ck('later-era limitation remains explicit','No 2020–2023-only model or exposure-by-period interaction was available.'in main)
 ck('weather exclusion is not inferred','the weather models do not exclude a weather contribution'in main)
 ck('interpretation is exploratory','not evidence that cold became protective'in main)

 freeze=json.loads((M/'frozen_evidence_manifest.json').read_text())
 ck('all immutable evidence hashes equal the selected PR103 revision',all(hashlib.sha256((ROOT/r['path']).read_bytes()).hexdigest()==r['sha256']for r in freeze['sources']))
 ck('manuscript architecture follows the selected publication register',[k for k in re.findall(r'^## (.+)$',main,re.M)]==['Abstract','Introduction','Results','Discussion','Methods','Data availability','Code availability','References'])
 ck('all numerical tables retained in supplement',not mt and len(st)==9)
 fit=[r for r in read(SOURCES[4])if r['family']=='negative_binomial'and r['offset_policy']=='days_only'and r['model_structure']=='single_exposure']
 ck('all twelve released diagnostic rows match source',st[8]==[[r['outcome'].upper(),dict(zip(EXPS,['Mean temperature','Mean maximum','Mean minimum','Hot nights','Very hot days','Cold days']))[r['exposure']],f"{float(r['theta']):.3f}",f"{float(r['dispersion']):.3f}",f"{float(r['residual_acf1']):.3f}",f"{float(r['ljung_box_p_lag6']):.3g}"]for r in fit])
 abstract=main.split('## Abstract\n\n')[1].split('\n\n**Keywords')[0]
 body=main.split('## Introduction\n')[1].split('## Data availability')[0]
 body=re.sub(r'\*\*Figure .*?(?=\n\n)','',body,flags=re.S);body=re.sub(r'!\[.*?\]\(.*?\)','',body);body=re.sub(r'^#+.*$','',body,flags=re.M)
 ck('abstract and prose meet editorial length targets',180<=len(abstract.split())<=220 and 2500<=len(body.split())<=3000)
 ck('figure and supplement cross references are complete',set(re.findall(r'^\*\*Figure (\d+)\.',main,re.M))=={'1','2','3'} and set(re.findall(r'^\*\*Supplementary Table S(\d+)\.',supp,re.M))==set(map(str,range(1,8))) and set(re.findall(r'^\*\*Supplementary Figure S(\d+)\.',supp,re.M))=={'1','2'})
 out={'date':'2026-10-04','status':'PASS','scope':'Direct source-to-table/value and interpretation checks; no health models run; no governance approval inferred','checks':checks,'tables':{'annual':st[0],'official_nested_full':st[3],'disjoint_periods':st[1],'full_panel':st[2],'calendar_sensitivity':st[4:6],'uncertainty':st[6:8],'diagnostics':st[8]},'nested_continuous':{o:{e:ci(s[o,e,'pre_covid'])for e in EXPS[:3]}for o in ['chd','hf']},'sources':{f:hashlib.sha256((ROOT/f).read_bytes()).hexdigest()for f in SOURCES},'freeze_anchor_commit':freeze['anchor_commit'],'word_counts':{'abstract':len(abstract.split()),'main_prose':len(body.split())},'claims':{'C01':{'claim':'annual trajectory and totals','source':SOURCES[0],'transform':'sum/annual totals; percent difference against2019'},'C02':{'claim':'disjoint means are descriptive','source':SOURCES[1]},'C03':{'claim':'complete twelve-fit family and multiplicity','source':SOURCES[2],'filter':'negative_binomial/days_only/single_exposure/NW6'},'C04':{'claim':'calendar and overlapping-window sensitivity','source':SOURCES[3]},'C05':{'claim':'uncertainty constructions','source':SOURCES[2]},'C06':{'claim':'released residual/dispersion diagnostics','source':SOURCES[4]},'C07':{'claim':'recorded outcome, absent cause/person-time','source':'reports/data_receipt_2026-08-07.md'},'C08':{'claim':'external clinical context, not local measurement','source':'literature/covid_novelty_sources_2026-10-02.json','limits':'qualitative evidence only; access statuses retained in source audit'}},'manuscript_sha256':hashlib.sha256((M/'Manuscript_covid_period_draft.md').read_bytes()).hexdigest(),'figures':fm['figures'],'literature_register_sha256':hashlib.sha256((ROOT/'literature/covid_novelty_sources_2026-10-02.json').read_bytes()).hexdigest(),'supplement_sha256':hashlib.sha256((M/'Supplement_covid_period.md').read_bytes()).hexdigest()}
 (M/'claim_ledger.yml').write_text('# JSON is a YAML subset. Audit of the current COVID revision; historical ledger is in Git.\n'+json.dumps(out,indent=2,ensure_ascii=False)+'\n')
 print(f"PASS: {len(checks)} source and claim checks")

def build():
 from docx import Document
 from docx.shared import Pt,RGBColor,Inches
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 for source,target,header in [('Manuscript_covid_period_draft.md','Heat_CVD_Manuscript_covid_period.docx','Pandemic era first hospitalisation counts'),('Supplement_covid_period.md','Supplement_covid_period.docx','Supplementary results')]:
  b,parts=blocks((M/source).read_text());doc=Document();b.configure_document(doc,header,body_size=11.5)
  doc.sections[0].left_margin=Inches(.75);doc.sections[0].right_margin=Inches(.75)
  doc.sections[0].header.paragraphs[0].clear()
  footer=doc.sections[0].footer.paragraphs[0]
  footer.alignment=1
  field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='1';run.append(text);field.append(run);footer._p.append(field)
  for style in ['Title','Heading 1','Heading 2','Heading 3']:
   doc.styles[style].font.name='Times New Roman';doc.styles[style].font.color.rgb=RGBColor(0,0,0)
   for border in list(doc.styles[style]._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
  for kind,lines in parts:
   txt=' '.join(lines)
   if kind=='h0':
    p=doc.add_paragraph(style='Title');p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(10);b.add_rich_runs(p,txt,size=16);continue
   if kind in ['h1','h2','h3']:
    p,_=b.add_heading(doc,txt,1 if kind=='h1'else 2,size=11.5);p.style='Heading 1'if kind=='h1'else'Heading 2';continue
   if kind=='equation':
    p=doc.add_paragraph();p.alignment=1;p.paragraph_format.keep_with_next=True
    def mr(text):
     r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r
    def sub(base,index):
     x=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(base));u=OxmlElement('m:sub');u.append(mr(index));x.extend([e,u]);return x
    math=OxmlElement('m:oMath');math.extend([mr('log E('),sub('Y','t'),mr(') = log('),sub('d','t'),mr(') + α + β'),sub('X','t'),mr(' + ')])
    summ=OxmlElement('m:sSubSup');e=OxmlElement('m:e');e.append(mr('Σ'));lo=OxmlElement('m:sub');lo.append(mr('m = 2'));hi=OxmlElement('m:sup');hi.append(mr('12'));summ.extend([e,lo,hi]);math.append(summ)
    math.extend([sub('γ','m'),mr(' I('),sub('month','t'),mr(' = m) + s(t; 4 df).   (1)')]);p._p.append(math);continue
   if kind=='table':
    headers,rows=b.parse_table(lines);n=len(headers)
    widths=([.85,2.825,2.825]if n==3 else[1.75,1.58,1.58,1.59]if n==4 else[.65,1.55,2.4,.85,1.05]if n==5 and headers[0]=='Outcome'else[1.8,1.175,1.175,1.175,1.175]if n==5 else[.65,1.6,1.0,1.1,1.0,1.15])
    if headers[0]=='Period':widths=[1.55,.65,2.15,2.15]
    if headers[0]=='Outcome' and n==4:widths=[.6,1.65,2.125,2.125]
    b.add_table(doc,headers,rows,font_size=9.5,widths=widths)
    table=doc.tables[-1];border=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
     node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');border.append(node)
    table._tbl.tblPr.append(border);continue
   if kind=='image':
    match=re.search(r'\]\(([^)]+)\)',txt);b.add_picture(doc,(M/match.group(1)).resolve(),width=6.5,max_height=5.4);continue
   plain=b.clean_cell(txt)
   if plain.startswith(('Table ','Figure ','Supplementary Table ','Supplementary Figure ')):
    b.add_caption(doc,plain,size=10.5);continue
   if re.match(r'^\d+\. ',plain):
    b.add_body(doc,txt,size=10.5,first_line=False,align='left');continue
   b.add_body(doc,txt,size=11.5,first_line=False,align='left')
  doc.core_properties.title=parts[0][1][0];doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.save(M/target);print('Wrote',target)
if __name__=='__main__':
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('--audit',action='store_true');a=p.parse_args();audit()if a.audit else build()
