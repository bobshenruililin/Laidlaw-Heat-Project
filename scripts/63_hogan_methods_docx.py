#!/usr/bin/env python3
"""Methods-only Word paste for Hogan's live file (24 Aug 2026).

Do not email this file. Paste into the shared document.
Weather paragraph matches Hogan's 24 August PDF wording.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path("/workspace")
SRC = ROOT / "manuscript/live_collaborative/Heat_CVD_Manuscript_live_update.md"
OUT = ROOT / "manuscript/live_collaborative/Heat_CVD_Methods_20260824_paste.docx"


def set_run_font(run, *, italic=False, bold=False, size=12):
    run.italic = italic
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)


def add_para(doc, text, *, bold=False, italic=False, size=12, first_line=True, space_after=8):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = Cm(1.27 if first_line else 0)
    run = p.add_run(text)
    set_run_font(run, italic=italic, bold=bold, size=size)
    return p


def main() -> None:
    md = SRC.read_text()
    block = md.split("## Methods\n", 1)[1].split("## Results\n", 1)[0].strip()
    doc = Document()
    add_para(doc, "Methods — paste into Hogan live file (24 August 2026)", bold=True, first_line=False, size=14)
    add_para(
        doc,
        "Replace the live Methods with the text below. Do not email this file. "
        "The weather paragraph is Hogan’s 24 August wording.",
        italic=True,
        first_line=False,
        size=10,
    )
    for raw in block.split("\n"):
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("### "):
            add_para(doc, line[4:], bold=True, first_line=False, size=12, space_after=6)
        elif line.startswith("#### "):
            add_para(doc, line[5:], bold=True, italic=True, first_line=False, size=12, space_after=4)
        elif line.startswith("\\[") or line == "\\]" or line.startswith("\\]"):
            continue
        elif "log \\mathrm{E}" in line or line.startswith("\\log"):
            add_para(
                doc,
                "log E(Y_t) = log(d_t) + alpha + beta X_t + calendar-month indicators + s(t; 4 df).",
                italic=True,
                first_line=False,
                size=11,
            )
        else:
            add_para(doc, line.lstrip("# "))
    # equation as a standalone line from the markdown math
    # (already included as surrounding sentences)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
