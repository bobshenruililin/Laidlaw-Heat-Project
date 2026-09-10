"""Contract tests for the 10 September analysis-window live package.

No governed panel is read and no health model is fitted.
"""
from __future__ import annotations

import hashlib
from pathlib import Path

import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
LIVE = ROOT / "manuscript/live_collaborative/Heat_CVD_Manuscript_live_update.md"
DOCX = ROOT / "manuscript/live_collaborative/Heat_CVD_Manuscript_20260910_hogan.docx"
PDF = ROOT / "manuscript/live_collaborative/Heat_CVD_Manuscript_20260910_hogan.pdf"
SUPP_MD = ROOT / "manuscript/live_collaborative/supplement_live_track.md"
SUPP_DOCX = ROOT / "manuscript/live_collaborative/Heat_CVD_Supplement_20260910_hogan.docx"
SUPP_PDF = ROOT / "manuscript/live_collaborative/Heat_CVD_Supplement_20260910_hogan.pdf"
ARCHIVE = ROOT / "manuscript/archive/thermal_extremes_2026-08"

HOGAN_OPEN = "Meteorological data was obtained from the HKO."
HOGAN_AVG = (
    "All monthly data were derived by taking the average of daily data in each calendar month."
)


def _norm(text: str) -> str:
    return " ".join(text.replace("\u00ad", "").split()).replace("–", "-").replace("—", "-")


def _pages(path: Path) -> list[str]:
    with pymupdf.open(path) as pdf:
        return [_norm(page.get_text()) for page in pdf]


def test_live_scientific_body_is_nested_window_not_prepost_effect():
    text = LIVE.read_text(encoding="utf-8")
    low = text.lower()
    assert text.startswith("# Analysis-window sensitivity")
    assert text.count(HOGAN_OPEN) == 1
    assert text.count(HOGAN_AVG) == 1
    assert "nested fits do not estimate a pre/post effect" in low
    assert "### nested-window official-day panel" in low
    assert "no interaction between thermal exposure and period was fitted" in low
    assert "84 pre-2020 months are contained in the full window" in low
    assert "interval overlap is not a test of a window difference" in low
    assert "not evidence that weather was irrelevant" in low
    assert "do not show improved cardiovascular health" in low
    assert "all twelve full-window *q*-values exceeded 0.19" in low
    for forbidden in (
        "health improved",
        "gate 3",
        "hogan asked",
        "candidate article",
        "admissions averted",
        "failing left ventricle",
    ):
        assert forbidden not in low


def test_word_contract_matches_hogan_a4_times_and_comments():
    assert DOCX.is_file() and DOCX.stat().st_size > 100_000
    doc = Document(str(DOCX))
    section = doc.sections[0]
    assert round(section.page_width.mm, 1) == 210.0
    assert round(section.page_height.mm, 1) == 297.0
    for margin in (
        section.top_margin,
        section.bottom_margin,
        section.left_margin,
        section.right_margin,
    ):
        assert round(margin.cm, 2) == 2.54
    text = "\n".join(p.text for p in doc.paragraphs)
    assert HOGAN_OPEN in text and HOGAN_AVG in text
    assert "Analysis-window sensitivity" in text
    assert "I(count/5)" in text
    assert len(doc.tables) == 3
    assert len([r for r in doc.part.rels.values() if "image" in r.reltype]) == 3
    comments = list(doc.comments)
    assert len(comments) >= 7
    joined = " ".join(c.text for c in comments).lower()
    assert "irb" in joined
    assert "roro" in joined
    assert "hogan" in joined
    assert "nested" in joined
    for forbidden in (
        "paste",
        "circulate",
        "shared live",
        "pipeline",
        "gate 3",
        "core panel",
        "rscript",
    ):
        assert forbidden not in joined


def test_main_pdf_has_distinct_float_pages_and_no_process_banner():
    assert PDF.is_file() and PDF.stat().st_size > 100_000
    pages = _pages(PDF)
    full = "\n".join(pages)
    assert len(pages) >= 18
    for required in (
        "Analysis-window sensitivity",
        HOGAN_OPEN,
        HOGAN_AVG,
        "1.022 (1.002-1.042)",
        "1.073 (1.006-1.144)",
        "Table 1. Outcome summary",
        "Table 2. Official-day count ratios",
        "Table 3. Uncertainty ladder",
        "Figure 1. First-event",
        "Figure 2. Official cold days",
        "Figure 3. Model 1 count ratios",
        "Supplementary Table S9",
        "UW XX-XXX",
        "None.",
    ):
        assert required in full, required
    for forbidden in ("Gate 3", "candidate article", "Hogan asked"):
        assert forbidden.lower() not in full.lower()

    def page_with(needle: str) -> int:
        hits = [idx for idx, text in enumerate(pages) if needle in text]
        assert hits, needle
        return hits[0]

    floats = (
        page_with("Table 1. Outcome summary"),
        page_with("Table 2. Official-day count ratios"),
        page_with("Table 3. Uncertainty ladder"),
        page_with("Figure 1. First-event"),
        page_with("Figure 2. Official cold days"),
        page_with("Figure 3. Model 1 count ratios"),
    )
    assert len(set(floats)) == 6
    assert "HF mean minimum temperature / 1 °C" in pages[floats[2]]


def test_supplement_resolves_s1_s7_s9_s10_and_is_formatted():
    text = SUPP_MD.read_text(encoding="utf-8")
    assert SUPP_DOCX.is_file() and SUPP_DOCX.stat().st_size > 100_000
    assert SUPP_PDF.is_file() and SUPP_PDF.stat().st_size > 100_000
    pages = _pages(SUPP_PDF)
    full = "\n".join(pages)
    for required in (
        "Supplementary Table S1",
        "Supplementary Table S7",
        "Supplementary Table S8",
        "Supplementary Table S9",
        "Supplementary Table S10",
        "Supplementary Note S1",
        "SYNTHETIC_CALIBRATION",
        "Later years were hotter",
        "Figure S5 remains unassigned",
    ):
        assert required in text
        assert required in full
    assert "health improved" not in full.lower()


def test_thermal_archive_core_hashes_remain_frozen():
    snapshot = ARCHIVE / "live_collaborative_snapshot"
    expected = {
        "Heat_CVD_Manuscript_live_update.md": "21c024bda975f9f5",
        "Heat_CVD_Manuscript_20260824_hogan.docx": "852702ded2c4e968",
        "Heat_CVD_Manuscript_20260824_hogan.pdf": "96b69b7b7ba38513",
        "claim_ledger.yml": "28a8cf48fd5b3268",
    }
    for name, want in expected.items():
        got = hashlib.sha256((snapshot / name).read_bytes()).hexdigest()[:16]
        assert got == want, (name, got, want)
