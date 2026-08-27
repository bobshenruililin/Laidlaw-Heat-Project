#!/usr/bin/env python3
"""Build the 15 August paste Word file for Hogan's live collaborative manuscript.

Do not email this file. Paste into the shared live document.

24 August 2026: this builder is provenance for the 15 August full-manuscript
Word copy. It is NOT the Methods paste source. Hogan's 24 August Methods live
in Heat_CVD_Manuscript_live_update.md and Heat_CVD_Methods_20260824_paste.docx
(scripts/63_hogan_methods_docx.py). Do not use this script to replace those Methods.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement

ROOT = Path("/workspace")
OUT = ROOT / "manuscript/live_collaborative/Heat_CVD_Manuscript_20260815_collab_draft.docx"
FIGS = {
    1: ROOT / "figures/live_identification/figure_B_first_event_depletion.png",
    2: ROOT / "figures/live_identification/figure_A_cold_day_identification.png",
    3: ROOT / "figures/live_identification/figure_D_trend_depletion_sensitivity.png",
}

HOGAN_WEATHER = (
    "Meteorological data was obtained from the HKO. The monthly variables included: "
    "temperature (mean, minimum, maximum), relative humidity, total rainfall, "
    "the number of hot nights (Tmin ≥ 28°C), the number of very hot days (Tmax ≥ 33°C), "
    "the number of extremely hot days (Tmax ≥ 35°C), and the number of cold days (Tmin ≤ 12°C). "
    "Monthly mean pollutant levels for nitrogen dioxide (NO2), sulfur dioxide (SO2), and ozone (O3) "
    "were obtained from the Hong Kong Environmental Protection Department (EPD). "
    "All monthly data were derived by taking the average of daily data in each calendar month."
)


def set_run_font(run, *, italic=False, bold=False, size=12, color=None, superscript=False, subscript=False):
    run.italic = italic
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    run.font.superscript = superscript
    run.font.subscript = subscript


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    set_run_font(run, size=10, italic=True)


def set_paragraph_format(p, *, first_line=True, space_after=8, space_before=0, align="left", size=12):
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if first_line:
        pf.first_line_indent = Cm(1.27)
    else:
        pf.first_line_indent = Cm(0)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, space_before=12 if level == 1 else 8, space_after=6)
    run = p.add_run(text)
    set_run_font(run, bold=True, size=14 if level == 1 else 12)
    return p


def body(doc, text, *, first_line=True, italic=False, bold=False, size=12, space_after=8):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=first_line, space_after=space_after, align="justify")
    run = p.add_run(text)
    set_run_font(run, italic=italic, bold=bold, size=size)
    return p


def mixed_body(doc, parts, *, first_line=True, space_after=8, align="justify"):
    """parts: list of (text, kwargs for set_run_font)."""
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=first_line, space_after=space_after, align=align)
    for text, kw in parts:
        run = p.add_run(text)
        set_run_font(run, **kw)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, space_before=6, space_after=10)
    run = p.add_run(text)
    set_run_font(run, italic=True, size=11)
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        set_paragraph_format(p, first_line=False, space_after=2, space_before=2)
        run = p.add_run(h)
        set_run_font(run, bold=True, size=10)
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            p = cells[c_i].paragraphs[0]
            set_paragraph_format(p, first_line=False, space_after=1, space_before=1)
            run = p.add_run(str(val))
            set_run_font(run, size=10)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()


def add_picture(doc, path, width_in=6.2):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, space_after=4, align="center")
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.clear()
    set_paragraph_format(hp, first_line=False, space_after=0)
    r = hp.add_run("Thermal extremes and first CHD/HF hospitalisation")
    set_run_font(r, italic=True, size=10, color=RGBColor(0x55, 0x55, 0x55))
    r2 = hp.add_run("\t")
    set_run_font(r2, size=10)
    add_page_number(hp)
    # tab stop right for page number
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), "9360")
    tabs.append(tab)
    hp._p.get_or_add_pPr().append(tabs)

    note = doc.add_paragraph()
    set_paragraph_format(note, first_line=False, space_after=10)
    nr = note.add_run(
        "Paste source for Hogan’s shared live file (15 August 2026). "
        "Do not email this as a parallel manuscript. "
        "The weather paragraph is Hogan’s 28 July text, copied verbatim. "
        "Reply to his comment threads from hogan_comment_paste_replies.md."
    )
    set_run_font(nr, italic=True, size=10, color=RGBColor(0x66, 0x33, 0x00))

    t = doc.add_paragraph()
    set_paragraph_format(t, first_line=False, space_after=6, align="center")
    tr = t.add_run(
        "Thermal extremes and first hospitalisation after coronary heart disease or heart failure diagnosis in Hong Kong, 2013–2023"
    )
    set_run_font(tr, bold=True, size=16)

    rt = doc.add_paragraph()
    set_paragraph_format(rt, first_line=False, space_after=10, align="center")
    rtr = rt.add_run("Running title: Thermal extremes and first CHD/HF hospitalisation")
    set_run_font(rtr, italic=True, size=11)

    auth = doc.add_paragraph()
    set_paragraph_format(auth, first_line=False, space_after=2, align="center")
    for text, kw in [
        ("Bob Ruililin Shen", {}),
        ("1", {"superscript": True}),
        (", Author 2", {}),
        ("1", {"superscript": True}),
        (", Author 3", {}),
        ("1", {"superscript": True}),
        (", Author 4", {}),
        ("1", {"superscript": True}),
        (", David Makram Bishai", {}),
        ("1", {"superscript": True}),
    ]:
        run = auth.add_run(text)
        set_run_font(run, size=12, **kw)

    aff = doc.add_paragraph()
    set_paragraph_format(aff, first_line=False, space_after=8, align="center")
    a1 = aff.add_run("1")
    set_run_font(a1, size=10, superscript=True)
    a2 = aff.add_run(
        " School of Public Health, Li Ka Shing Faculty of Medicine, The University of Hong Kong, Hong Kong SAR, China"
    )
    set_run_font(a2, size=10, italic=True)

    body(
        doc,
        "Author order after the first and last positions remains to be confirmed with the supervising investigator. Scientific roles are recorded in the Acknowledgements.",
        first_line=False,
        italic=True,
        size=10,
    )

    heading(doc, "Abstract")
    mixed_body(
        doc,
        [
            ("Background. ", {"bold": True}),
            (
                "Hong Kong records more hot nights than a decade ago, while cold days persist. Whether the cold-dominant pattern of earlier daily cardiac-admission studies extends to monthly first-hospitalisation counts is untested.",
                {},
            ),
        ],
    )
    mixed_body(
        doc,
        [
            ("Methods. ", {"bold": True}),
            (
                "We analysed 132 territory-months (2013–2023) of monthly counts of first hospitalisation after a first diagnosis of coronary heart disease (CHD; 156,156 events) or heart failure (HF; 29,681 events) among people with type 2 diabetes and/or hypertension. Admission cause was not recorded. Three continuous temperature measures and official counts of hot nights, very hot days, and cold days each entered a separate negative-binomial model with calendar-month indicators, a 4-df time spline, and a days-in-month offset. Intervals used model-based, HC1, and Newey–West lag-3 and lag-6 constructions. Benjamini–Hochberg q-values covered the twelve exploratory contrasts. A method for recovering daily exposure coefficients from monthly sums failed simulation calibration and was not applied to the health series.",
                {},
            ),
        ],
    )
    mixed_body(
        doc,
        [
            ("Results. ", {"bold": True}),
            (
                "Under Newey–West lag-6 reporting, the count ratio was 1.022 (1.002–1.042) per five hot nights for CHD and 1.073 (1.006–1.144) per five cold days for HF, both with q = 0.192; all twelve q-values exceeded 0.19. For CHD hot nights, model-based and HC1 intervals included 1, as did the pre-2020 estimate (1.011, 0.991–1.032). For HF cold days, all four constructions excluded 1, and the pre-2020 estimate was stronger (1.113, 1.053–1.176). Official cold days fell almost entirely in December–February (141 of 145), identifying that contrast between winters.",
                {},
            ),
        ],
    )
    mixed_body(
        doc,
        [
            ("Conclusions. ", {"bold": True}),
            (
                "The data do not support a multiplicity-protected differential thermal claim. The HF cold-day association is concordant across standard-error methods; the CHD hot-night association depends on the construction and on including 2020–2023. Estimates are ecological monthly count ratios for a first-event outcome without recorded admission cause.",
                {},
            ),
        ],
    )
    mixed_body(
        doc,
        [
            ("Keywords: ", {"bold": True}),
            (
                "hot nights; cold days; coronary heart disease; heart failure; Hong Kong; monthly time series",
                {},
            ),
        ],
        first_line=False,
    )

    heading(doc, "Introduction")
    body(
        doc,
        "Temperature-related cardiovascular morbidity remains a concern in subtropical cities, where intense humid heat coexists with episodic cold [9]. Hong Kong is a dense, ageing city whose thermal profile shifted within a single decade. At the Hong Kong Observatory (HKO), hot nights increased from 10 in 2013 to 61 in 2021, and very hot days from 17 to 54 in the same comparison [2,4]. Cold days persisted across the period: 14 in 2013, 13 in 2021, and 14 in 2023 [2,4,6]. A contemporary local analysis therefore has to hold heat and cold in one design.",
    )
    body(
        doc,
        "Local evidence on cardiac admissions is daily, diagnosis-specific, and drawn from earlier decades. Using data from 2000–2009, Goggins et al. reported that a 1 °C decrease below an estimated 24 °C threshold was associated with a 3.7% increase in acute myocardial infarction admissions, and found no statistically significant heat association in Hong Kong, Taipei, or Kaohsiung [1]. Same-day nitrogen dioxide was the strongest pollutant predictor in that study [1]. An analysis of public-hospital stroke admissions during 1999–2006 found an inverse association of haemorrhagic stroke with temperature, and a weaker ischaemic-stroke association below about 22 °C [11]. Closest to the present outcomes, Goggins and Chan analysed daily public-hospital heart-failure admissions and deaths during 2002–2011 [21]. They reported a cumulative relative risk of 2.63 (2.43–2.84) for admissions comparing 11 °C with 25 °C, over lags extending to 23 days [21]. These studies fix the local historical pattern. They estimate daily risk for cause-coded admissions, so their coefficients do not transfer to monthly counts of first hospitalisation.",
    )
    body(
        doc,
        "Nighttime heat can be encoded in more than one way. Guo et al. analysed daily unplanned emergency hospitalisations in Hong Kong during the hot seasons of 2000–2019 [17]. After adjustment for multi-day mean temperature, the official hot-night indicator (minimum temperature ≥ 28 °C) showed no overall association with non-cancer, non-external hospitalisation over lags 0–4 days (excess relative risk −0.2%, 95% CI −1.2% to 0.7%) [17]. An hourly excess-heat metric for 20:00–07:59 was associated with higher hospitalisation, including a 3.1% (1.5–4.8%) increase at its extreme value [17]. Official monthly counts of hot nights are, by construction, a coarser encoding than nighttime intensity. They remain an interpretable public climatological series, and they are not a test of the hourly metric. Guo et al. did not study first CHD or HF hospitalisation in a diabetes or hypertension cohort [17].",
    )
    mixed_body(
        doc,
        [
            (
                "Air pollution also changed during the present study window. At general monitoring stations, mean nitrogen dioxide declined from 53.7 µg m",
                {},
            ),
            ("−3", {"superscript": True}),
            (" in 2013 to 32.1 µg m", {}),
            ("−3", {"superscript": True}),
            (" in 2023, and fine particulate matter (PM2.5) from 30.8 to 14.6 µg m", {}),
            ("−3", {"superscript": True}),
            (" [7]. Ozone rose over the same years, from 42.6 to 58.3 µg m", {}),
            ("−3", {"superscript": True}),
            (" [7]. These concentrations describe 2013–2023 and are not a reconstruction of the 2000–2009 window analysed by Goggins et al. [1]. Temperature and pollutants are not interchangeable covariates. Ozone is coupled with hot, sunny conditions, and may confound the thermal association or lie on its pathway.", {}),
        ],
    )
    body(
        doc,
        "Daily mortality studies in Hong Kong answer a further, distinct question. Liu et al. (2020) reported a cold-dominant attributable fraction for cause-specific mortality during 2006–2016 (4.72% for cold versus 0.16% for heat), and a larger fraction for moderate non-optimum temperatures than for extremes (4.25% versus 0.63%) [12]. Liu et al. (2026) estimated 1,455 to 3,238 model-based excess heat deaths for 2014–2023, the range reflecting which local heatwave definition was applied [13]. Attributable fractions and modelled excess deaths are mortality burden quantities. They cannot be rescaled into monthly morbidity count ratios.",
    )
    body(
        doc,
        "This study estimates associations between specified monthly thermal-exposure contrasts and monthly counts of first hospitalisation after first CHD or HF diagnosis among people with type 2 diabetes and/or hypertension in Hong Kong from January 2013 through December 2023, carrying heat and cold in parallel and reporting the complete twelve-contrast panel in place of a single protected claim.",
    )

    heading(doc, "Methods")
    heading(doc, "Study design", level=2)
    body(
        doc,
        "The design is an ecological territory-month time series covering January 2013 through December 2023 (132 months). For each single-exposure model, the target quantity is a count ratio associated with a 1 °C or five-day exposure contrast, conditional on calendar month and a smooth function of time. Individual causal effects, principal-diagnosis CHD or HF, acute myocardial infarction, and daily triggering are not identified [14]. Monthly counts of cohort members still at risk of a first event were unavailable, so the offset is days in month and the estimates are monthly count ratios rather than cohort incidence-rate ratios [14]. The twelve-contrast model set was specified after the outcome series were available and is exploratory. We refer to the twelve separate single-exposure models as the core panel.",
    )

    heading(doc, "Health data", level=2)
    body(
        doc,
        "Governed monthly aggregates were constructed and released by the outcome co-investigator. The cohort comprised people diagnosed with type 2 diabetes and/or hypertension during 2013–2023. For each outcome, the event was the first recorded hospitalisation after the patient’s first diagnosis record for CHD or for HF. Diagnosis-record construction was assisted by a clinical collaborator. Admission cause was not recorded. Delivered columns were labelled chd_inpatient and hf_inpatient at territory × calendar-month grain. Inpatient-only semantics are inferred from those labels and remain subject to confirmation. Age–sex strata were not included. A stroke series was named in correspondence but was not attached, and no stroke analysis is reported. Field-level ICD lists and the executable timing rule remain with the outcome co-investigator and should replace this paragraph where they differ from it.",
    )

    heading(doc, "Weather and pollutants data", level=2)
    weather = body(doc, HOGAN_WEATHER)
    # mark Hogan paragraph with a light comment in the following italic note
    body(
        doc,
        "[Hogan weather paragraph above is copied verbatim from the 28 July live file, including “Meteorological data was obtained.” Do not edit it.]",
        first_line=False,
        italic=True,
        size=10,
        space_after=8,
    )
    body(
        doc,
        "Daily Headquarters series were aggregated to calendar months for analysis. Core exposures were monthly mean temperature, the monthly means of daily maximum and minimum temperature, and official counts of hot nights, very hot days, and cold days. The three counts were scaled per five days, so that a coefficient corresponds to a five-day contrast rather than to a single extreme day. Official cold days were concentrated in December–February (Results, Figure 2), so with calendar-month indicators in the model the remaining cold-day variation is a difference between years within winter months. A winter-only restriction is a labelled sensitivity, and not a remedy for that identification. Extremely hot days, relative humidity, rainfall, and sulfur dioxide were assembled with the weather series but did not enter the core panel. Fine particulate matter (PM2.5) from general EPD stations was assembled for staged pollution sensitivities. A station-month contributed to a pollutant mean only when at least 75% of expected observations were available, and missing measurements were not coded as zero. Territory-wide monthly concentrations were unweighted means across eligible general stations. Roadside stations were reserved. Nitrogen dioxide and particulate matter were specified to enter before ozone, because ozone is coupled with hot, sunny conditions. Staged pollution models exist in the archive; they are not adjusted versions of the core panel, and they were not used to claim that confounding is resolved. Operational heatwave-spell and hot-month rules beyond the official daily thresholds remain with the weather co-investigator, and provisional monthly-tail indicators were not used as confirmatory exposures.",
    )

    heading(doc, "Influenza", level=2)
    body(
        doc,
        "Monthly influenza activity was obtained from Centre for Health Protection Flu Express summaries. Coverage was 121 of 132 months, with the gap in early 2013. Missing influenza months were left missing and were never coded as zero. Influenza did not enter the core panel.",
    )

    heading(doc, "Population denominators and offsets", level=2)
    body(
        doc,
        "Age- and sex-specific mid-year population estimates from the Census and Statistics Department (Table 110-01001) were linearly interpolated to calendar-month midpoints [8]. Those denominators describe the general population, and not the diabetes or hypertension cohort still at risk of a first CHD or HF hospitalisation. Core models therefore used a days-in-month offset, log(dt), and are interpreted as monthly count ratios [14]. A general-population × days offset was retained only as a sensitivity. It does not convert the estimand into cohort incidence.",
    )

    heading(doc, "Statistical analysis", level=2)
    body(
        doc,
        "Let Yt be the monthly count in month t. Core models used a negative-binomial likelihood [10]: log E(Yt) = log(dt) + α + β Xt + month indicators + s(t; 4 df), where dt is days in month, Xt is a single exposure, and s(t; 4 df) is a natural cubic spline of month index. Each core model entered one exposure at a time: mean temperature, mean maximum temperature, mean minimum temperature, hot nights, cold days, or very hot days, for CHD and for HF. That set defines the twelve core contrasts. Quasi-Poisson models were a family sensitivity. Negative binomial was preferred because monthly counts were overdispersed relative to Poisson.",
    )
    body(
        doc,
        "Core intervals used Newey–West standard errors with lag 6 [19], reported together with model-based, HC1, and Newey–West lag-3 intervals, treating the fitted negative-binomial score equations as estimating equations. Lag 6 was the core reporting choice because residual serial dependence is expected in monthly counts and is material for CHD. Robust intervals are not automatically wider than model-based intervals, and the four-construction ladder is reported for that reason. No interval was selected because it excluded 1.",
    )
    body(
        doc,
        "Joint models that entered maximum and minimum temperature together, or the three official extreme-day counts together, were retained as collinearity diagnostics. Variance inflation factors were computed after residualising on calendar month and the trend spline.",
    )
    body(
        doc,
        "Robustness analyses comprised alternative trend smooths (3, 6, and 8 df), year fixed effects, exclusion of the first 12 or 24 months, a pre-2020 window (January 2013–December 2019; 84 months), COVID-phase adjustment, lag-0, lag-1, and lag-2 month exposures, and exclusion of the month with largest Cook’s distance. COVID phases were pre-COVID (through January 2020), early COVID (February 2020–December 2021), fifth wave (January–April 2022), late 2022 (May–December 2022), and post-reopening (from January 2023). Pearson residual autocorrelation and Ljung–Box tests at lags 6 and 12 were inspected. Benjamini–Hochberg q-values were computed across the twelve core contrasts [20].",
    )
    body(
        doc,
        "Archive models that entered pollution, humidity, or influenza are not adjusted versions of the core panel. They were not used to claim that confounding is resolved.",
    )

    heading(doc, "Constrained daily-exposure recovery", level=2)
    body(
        doc,
        "A constrained method for recovering daily exposure coefficients from monthly sums was evaluated in simulation against Hong Kong weather, following the Basagaña–Ballester aggregation framework [15,16]. Implementation was from the published estimating equation. Public code without a repository licence was not copied. Synthetic daily outcomes were matched to the disclosure-minimised seasonality, scale, overdispersion, and residual dependence of the monthly series. Admission of any real daily coefficient required all frozen calibration criteria to pass, including Type I error, coverage, relative bias, and sign recovery. A 500-replicate calibration against Hong Kong weather failed those requirements. A subsequent worst-cell re-summary of the same fits did not reverse that decision. No daily health coefficient was produced. This is a project-specific calibration result, and it does not show that the Basagaña–Ballester method fails in general. A full daily distributed-lag non-linear model is not identified from 132 monthly sums.",
    )

    heading(doc, "Software", level=2)
    body(
        doc,
        "Analyses were conducted in R 4.3.3 (2024-02-29). Negative-binomial models used MASS::glm.nb (MASS 7.3-60.0.1). Newey–West and HC1 standard errors used sandwich 3.1.3. Natural cubic splines used splines::ns.",
    )

    heading(doc, "Ethics and governance", level=2)
    body(
        doc,
        "Analyses used governed Hospital Authority monthly aggregates supplied under existing collaborative arrangements. Source monthly counts and merged panels are not redistributed. Aggregation reduces re-identification risk, but does not grant external-publication authority. Written confirmation from the outcome and supervising investigators remains required before journal submission. Whether the institutional review protocol requires amendment for this work remains a decision for the supervising investigator.",
    )

    heading(doc, "Results")
    heading(doc, "Outcome series", level=2)
    body(
        doc,
        "CHD contributed 156,156 first recorded hospitalisations over 132 months, a mean of 1,183.0 per month. HF contributed 29,681, a mean of 224.9 per month (Table 1). Annual totals fell by about half from 2013 to 2023 (CHD 23,830 to 12,323; HF 4,336 to 2,296), with a further dip in 2020 (CHD 10,237; HF 1,964). Over the same years, the interpolated Census and Statistics Department population aged 35 years or older rose by 17%. Figure 1 indexes both series to 2013. The decline in first-event counts is compatible with depletion of the at-risk set in a 2013–2023 diagnosis window, together with pandemic changes in care-seeking. It is not a measure of falling incidence in the territory, and it does not quantify the share attributable to depletion without monthly still-at-risk person-time. Mean monthly HF counts were highest in January (287) and lowest in September (196). CHD showed a milder winter elevation (January 1,386; September 1,097). After calendar-month indicators, the thermal estimates ask whether a given January or July differs from a typical January or July, and not whether winter counts exceed summer counts.",
    )
    caption(doc, "Table 1. Outcome summary.")
    add_table(
        doc,
        ["Outcome", "Period", "Months", "Total events", "Mean per month", "Event construction"],
        [
            ["Coronary heart disease", "2013–2023", "132", "156,156", "1,183.0", "First recorded hospitalisation after first CHD diagnosis"],
            ["Heart failure", "2013–2023", "132", "29,681", "224.9", "First recorded hospitalisation after first HF diagnosis"],
        ],
    )
    add_picture(doc, FIGS[1])
    caption(
        doc,
        "Figure 1. First-event counts fell while the general population aged 35 years or older rose. Index = 1 in 2013. CHD 23,830 → 12,323; HF 4,336 → 2,296. Source: governed Hospital Authority aggregate annual totals; Census and Statistics Department mid-year population aged 35 years or older.",
    )

    heading(doc, "Exposure context", level=2)
    body(
        doc,
        "At HKO Headquarters, hot nights rose from 10 in 2013 to 61 in 2021 and 56 in 2023 [6]. The 2013 count was about seven days below the 1981–2010 normal [2]. The 2021 count was the highest annual number of hot nights in the Observatory record at that time [4]. Very hot days rose from 17 to 54 over the same 2013–2021 comparison, and were 54 in 2023 [6]. Cold days were 14 in 2013 and 14 in 2023, with 1 in 2019 and 13 in 2021 [6]. These are exposure counts, and not health effects. Figure 2 shows official cold days by year and month. Of 145 cold days in 2013–2023, 141 fell in December–February (December 40, January 54, February 47), and 4 fell in March. Twenty-nine of 132 months carried at least one official cold day. The year 2019 contributed a single cold day. After calendar-month indicators, the remaining cold-day variation is between-year winter variation, and not a summer-versus-winter contrast.",
    )
    add_picture(doc, FIGS[2])
    caption(
        doc,
        "Figure 2. Official cold days by year and month, Hong Kong Observatory Headquarters, 2013–2023. Source: Hong Kong Observatory official daily cold-day flags (Tmin ≤ 12 °C), aggregated to calendar months.",
    )

    heading(doc, "Core panel", level=2)
    body(
        doc,
        "Table 2 reports the twelve separate-exposure count ratios under the days-in-month offset and Newey–West lag-6 intervals.",
    )
    caption(doc, "Table 2. Separate negative-binomial models with days-in-month offset and Newey–West lag-6 intervals.")
    add_table(
        doc,
        ["Outcome", "Exposure contrast", "Count ratio (95% CI)", "p", "BH q"],
        [
            ["CHD", "Mean temperature / 1 °C", "0.993 (0.976–1.010)", "0.424", "0.726"],
            ["CHD", "Mean maximum temperature / 1 °C", "0.993 (0.979–1.007)", "0.321", "0.642"],
            ["CHD", "Mean minimum temperature / 1 °C", "0.994 (0.977–1.012)", "0.509", "0.764"],
            ["CHD", "Hot nights / 5 days", "1.022 (1.002–1.042)", "0.032", "0.192"],
            ["CHD", "Cold days / 5 days", "0.995 (0.949–1.043)", "0.823", "0.900"],
            ["CHD", "Very hot days / 5 days", "0.999 (0.974–1.025)", "0.962", "0.962"],
            ["HF", "Mean temperature / 1 °C", "0.974 (0.947–1.002)", "0.069", "0.207"],
            ["HF", "Mean maximum temperature / 1 °C", "0.981 (0.958–1.005)", "0.113", "0.272"],
            ["HF", "Mean minimum temperature / 1 °C", "0.973 (0.947–1.000)", "0.050", "0.202"],
            ["HF", "Hot nights / 5 days", "1.003 (0.976–1.031)", "0.825", "0.900"],
            ["HF", "Cold days / 5 days", "1.073 (1.006–1.144)", "0.031", "0.192"],
            ["HF", "Very hot days / 5 days", "0.995 (0.963–1.028)", "0.764", "0.900"],
        ],
    )
    body(
        doc,
        "All twelve q-values exceeded 0.19, with a minimum of 0.192. No contrast meets a multiplicity-protected confirmatory threshold in this exploratory panel. The HF mean-minimum-temperature interval narrowly included 1, with an unrounded upper bound of 1.00005.",
    )
    body(
        doc,
        "For CHD, the continuous monthly temperature estimates were near null, between 0.993 and 0.994, whereas the official hot-night count estimate was larger under Newey–West lag-6 reporting. These encodings describe different monthly thermal questions. The contrast does not imply that a hot-night effect was missed by the means, and it is not a test of hourly nighttime excess heat [17]. For HF, the continuous temperature estimates were weakly inverse, and the cold-day count estimate was positive.",
    )

    heading(doc, "Uncertainty ladder", level=2)
    caption(doc, "Table 3. Uncertainty ladder for the two leading exploratory contrasts.")
    add_table(
        doc,
        ["Outcome and exposure", "Model", "HC1", "NW3", "NW6"],
        [
            ["CHD hot nights / 5 days", "1.022 (0.995–1.049)", "1.022 (0.997–1.047)", "1.022 (1.0003–1.0439)", "1.022 (1.002–1.042)"],
            ["HF cold days / 5 days", "1.073 (1.023–1.125)", "1.073 (1.011–1.138)", "1.073 (1.007–1.143)", "1.073 (1.006–1.144)"],
        ],
    )
    body(
        doc,
        "For CHD hot nights, the model-based and HC1 intervals included 1, while both Newey–West intervals excluded 1 on the unrounded scale (lag 3: 1.000253 to 1.043860). The ladder narrowed from the model-based interval to the Newey–West intervals for that contrast. Robust intervals are not automatically wider, and the direction of that change is itself informative. For HF cold days, all four constructions excluded 1, and q remained 0.192. Concordance across standard-error methods does not create multiplicity protection.",
    )

    heading(doc, "Joint models and residual diagnostics", level=2)
    body(
        doc,
        "After month and trend residualisation, maximum and minimum temperature retained a variance inflation factor of 4.66 when entered jointly. Hot nights and very hot days retained variance inflation factors near 1.96. In the joint extreme-day model, the CHD hot-night count ratio was 1.045 (1.015–1.075), compared with 1.022 in the separate model. The HF cold-day estimate was the same in the joint and separate models, at 1.073. Joint coefficients are diagnostics, and not preferred estimates.",
    )
    mixed_body(
        doc,
        [
            (
                "Pearson residual autocorrelation at lag 1 was 0.508 in the CHD hot-night model and 0.146 in the HF cold-day model (Supplementary Figure S1). Ljung–Box tests at lag 6 gave ",
                {},
            ),
            ("p", {"italic": True}),
            (" < 10", {}),
            ("−7", {"superscript": True}),
            (" for every CHD core model and ", {}),
            ("p", {"italic": True}),
            (" > 0.3 for every HF core model. Serial dependence is part of the inferential problem for CHD, which is why the four-construction ladder is shown rather than a single interval.", {}),
        ],
    )

    heading(doc, "Sensitivity analyses", level=2)
    body(
        doc,
        "Offset choice changed the core count ratios only trivially. Across trend, window, and COVID-phase specifications, the CHD hot-night ratio ranged from 1.011 to 1.025, and the HF cold-day ratio from 1.043 to 1.113 (Figure 3). The HF cold-day association was strongest before 2020 (1.113, 1.053–1.176). The CHD hot-night association was weaker and compatible with 1 in the pre-2020 window (1.011, 0.991–1.032) and under COVID-phase adjustment (1.013, 0.994–1.033). The decline in first events and the rise in hot nights are both strong trends across this window, and the same 4-df spline absorbs both. Lag-1 month models attenuated the CHD hot-night association toward 1 (1.010, 0.979–1.042). The HF cold-day association remained elevated at lag 1 (1.073, 1.014–1.135) and was weaker at lag 2 (1.053, 0.985–1.127). Lag-one month is a labelled sensitivity, and it is not a daily lag curve. The months with largest Cook’s distance were February 2020 for CHD and February 2022 for HF cold days. Excluding the most influential month left both exploratory directions unchanged (CHD hot nights 1.021, 1.001–1.040; HF cold days 1.088, 1.034–1.144).",
    )
    add_picture(doc, FIGS[3])
    caption(
        doc,
        "Figure 3. Core count ratios across trend, window, and COVID-phase specifications (Newey–West lag-6 intervals). All twelve core contrasts are shown across nine specifications. The pre-2020 CHD hot-night interval includes 1, while the pre-2020 HF cold-day interval strengthens. Source: sensitivity estimates computed from the governed aggregate series (trend and depletion sensitivity table).",
    )
    body(
        doc,
        "The simulated daily-recovery method failed its worst-cell criteria: null Type I error ranged from 0.048 to 0.150, minimum coverage was 0.840, maximum non-null relative bias was 32.8, and the maximum moderate false-sign rate was 0.808. No real daily coefficient is reported. That refusal is specific to this monthly series and this implementation.",
    )

    heading(doc, "Discussion")
    for para in [
        "These data do not support a multiplicity-protected differential thermal claim for CHD relative to HF. Under Newey–West lag-6 reporting for the core panel, CHD first-hospitalisation counts were more closely associated with official hot-night burden than with mean temperature or cold days. HF counts were more closely associated with cold-day burden than with hot nights. Both patterns sit inside a twelve-contrast family in which every q-value exceeds 0.19. The complete panel is therefore reported, and no contrast is promoted to a primary result.",
        "The HF cold-day association is the more coherent of the two residual signals. It is concordant across all four standard-error constructions. It survives exclusion of the most influential pandemic month. It is not produced by entering correlated heat metrics jointly. It is strongest in the pre-2020 window, before pandemic disruption of care-seeking. It nevertheless remains unprotected by its q-value. Because official cold days fall almost entirely in December–February, and because only 29 of 132 months carry any official cold day, the association is identified from differences between winters rather than from a summer-versus-winter contrast.",
        "The CHD hot-night association is smaller, depends on the uncertainty method, and was not evident in the pre-2020 window. Model-based and HC1 intervals include 1, whereas the Newey–West intervals exclude 1 on the unrounded scale (lag 3: 1.000253 to 1.043860). Pearson residual autocorrelation at lag 1 is 0.508 in that model, so inference for the CHD series depends on how serial dependence is handled. For this contrast the Newey–West intervals were narrower than the model-based interval, which is atypical under positive residual autocorrelation; the exclusion of 1 therefore rests on the smaller robust variance estimate, and is a reason for caution rather than confirmation. No interval construction was chosen because it excluded 1. The estimate is also confined to specifications that include 2020–2023, which are the years in which hot nights peaked and in which care-seeking changed. Hot-night burden and the decline in first events both trend strongly across the window, and both are controlled by the same trend spline, so the hot-night contrast carries the harder identification problem of the two.",
        "The direction of the HF cold-day residual is consistent with earlier daily evidence that lower temperature was associated with higher heart-failure admissions in Hong Kong [21]. The comparison is between questions, and not between magnitudes. A cumulative daily relative risk comparing 11 °C with 25 °C is not a monthly count ratio per five official cold days, and the present event is a first hospitalisation after a first HF diagnosis without recorded admission cause. The CHD hot-night residual should likewise be read against Guo et al. rather than as a replication of it [17]. That study reported no overall association for the official hot-night flag after adjustment for mean temperature, and a positive association for hourly nighttime excess heat [17]. The present analysis uses monthly official counts, a first-event CHD series, and a later decade. A difference between monthly mean temperature and monthly official hot-night counts does not identify an intensity mechanism.",
        "Liu et al. (2020, 2026) remain complementary mortality baselines [12,13]. Their attributable fractions and excess-death totals cannot be rescaled into the present count ratios. The failed daily-recovery calibration is the corresponding methods limit: monthly sums do not automatically yield daily trigger estimates [15,16]. That refusal is specific to this series, this implementation, and this calibration standard. It is not evidence that recovery of daily effects from aggregated outcomes fails in general.",
        "Pollution and influenza are scientifically motivated in this setting [1,18]. Archive models that entered nitrogen dioxide, fine particulate matter, or ozone used a general-population offset and joint maximum and minimum temperature, so they are not adjusted versions of the separate core panel (Supplementary Table S9). On a monthly grain they cannot separate confounding from mediation for ozone, and no staged estimate was used to revise Table 2. On the 121 months with influenza data, an archive model associated higher influenza activity with higher CHD counts (Supplementary Table S7); that model is not an adjusted version of the core panel and does not adjust Table 2. Confounding by infection, ozone, or nitrogen dioxide is therefore unresolved. Absent cohort person-time, even a stable count ratio remains a count ratio [14].",
        "The contribution of this analysis is identification rather than estimation. A monthly aggregate series of first events can be reported in a way that shows where its information comes from. Four displays do that work here: which months carry the exposure contrast, whether the outcome decline is compatible with risk-set depletion, how far each interval moves across standard-error constructions, and which specifications each estimate depends on. A protected primary claim would have required a predeclared confirmatory contrast, multiplicity control that survives the twelve-contrast family, uncertainty constructions that are not chosen for null exclusion, and residual diagnostics that leave no substantial CHD serial correlation unaddressed. This analysis does not meet that bar. We report the complete panel and the refusals.",
    ]:
        body(doc, para)

    heading(doc, "Conclusion")
    body(
        doc,
        "Between 2013 and 2023, hot nights in Hong Kong increased while cold days persisted. In governed monthly aggregates for people with type 2 diabetes and/or hypertension, CHD first hospitalisations were more closely associated with hot nights, and HF first hospitalisations with cold days, than with the other thermal encodings examined. Neither association survived correction across the twelve comparisons. The CHD estimate additionally depended on the treatment of uncertainty and on the inclusion of 2020–2023. This analysis therefore contributes a set of hypotheses, and an explicit account of what monthly aggregate counts cannot settle. Better-denominated and more finely resolved data are required before a thermal effect on cardiac hospitalisation in this cohort can be estimated.",
    )

    heading(doc, "Strengths and limitations")
    mixed_body(
        doc,
        [
            ("Strengths. ", {"bold": True}),
            (
                "The twelve-contrast panel is reported in full, rather than filtered to its largest estimates. Uncertainty is shown as an explicit four-construction ladder rather than as a single interval. Extreme-day exposures use published official thresholds. Identification displays show the sources of cold-day variation, the first-event decline, and the dependence of each estimate on the analysis window. The daily-recovery analysis is reported as a refusal rather than as a coefficient.",
                {},
            ),
        ],
    )
    mixed_body(
        doc,
        [
            ("Limitations. ", {"bold": True}),
            (
                "Admission cause was not recorded, so an event is a first hospitalisation after a first diagnosis and not a cardiac-caused admission. Inpatient-only semantics are inferred from the delivered column labels and are not yet confirmed. Monthly counts of cohort members still at risk of a first event were unavailable, so the estimates are count ratios rather than incidence-rate ratios [14]. The design is ecological and monthly, so individual-level and daily-triggering interpretations are not identified [14]. Age, sex, and disease-subtype strata were not delivered. Residual serial correlation remains material for CHD, with lag-1 Pearson autocorrelation of 0.508 in the hot-night model. Official cold days are concentrated in December–February, and only 29 of 132 months carry any official cold day, so the HF cold-day estimate rests on differences between winters. The CHD hot-night estimate is not stable across analysis windows, and is compatible with 1 before 2020. Official hot-night counts are not hourly nighttime excess heat [17]. All exposures were measured at a single Observatory station and applied territory-wide, and within-territory exposure variation was not modelled. Reference rules for monthly hot-tail and cold-tail indicators are not yet fixed, and no such indicator was used as a confirmatory exposure. Confounding by pollution, humidity, and influenza is unresolved in the core panel. A stroke series was named in correspondence but was not delivered, and no stroke result is reported. External dissemination of the governed aggregates requires confirmation from the outcome and supervising investigators.",
                {},
            ),
        ],
    )

    heading(doc, "Declaration on the use of AI")
    body(
        doc,
        "AI-assisted tools were used for code scaffolding, methodological brainstorming, and review of potential data and modelling issues. Scientific claims, numerical results, and collaborator-owned Methods text were checked against governed outputs and the live weather paragraph.",
    )
    heading(doc, "Declaration of competing interest")
    body(doc, "The authors declare no competing interests.")
    heading(doc, "Acknowledgements")
    body(
        doc,
        "We thank the weather co-investigator for the live weather Methods, thermal-definition guidance, and comments on the Introduction; Zhenyuan Liu for constructing and releasing the governed Hospital Authority aggregates and for regression mentorship; Professor David Bishai for supervision; and Dr Jingjing Zhou for assistance with diagnosis-record construction. Author order and any co-first designation remain to be confirmed.",
    )
    heading(doc, "Data and code availability")
    body(
        doc,
        "A disclosure-minimised release of tables, figures, and machine-readable checks accompanies the analysis. Source monthly health counts and merged panels are governed and are not redistributed. External submission requires team confirmation.",
    )

    heading(doc, "References")
    refs = [
        "1. Goggins WB, Chan EYY, Yang CY. Weather, pollution, and acute myocardial infarction in Hong Kong and Taiwan. Int J Cardiol. 2013;168(1):243-249. doi:10.1016/j.ijcard.2012.09.087",
        "2. Hong Kong Observatory. The Year's Weather — 2013. https://www.hko.gov.hk/en/wxinfo/pastwx/ywx2013.htm",
        "3. Hong Kong Observatory. The Year's Weather — 2017. https://www.hko.gov.hk/en/wxinfo/pastwx/2017/ywx2017.htm",
        "4. Hong Kong Observatory. The Year's Weather — 2021. https://www.hko.gov.hk/en/wxinfo/pastwx/2021/ywx2021.htm",
        "5. Hong Kong Observatory. The Year's Weather — 2024. https://www.hko.gov.hk/en/wxinfo/pastwx/2024/ywx2024.htm",
        "6. Hong Kong Observatory. Hong Kong Observatory Open Data. https://www.hko.gov.hk/en/abouthko/opendata_intro.htm",
        "7. Environmental Protection Department. Environmental Protection Interactive Centre: Air Quality Data. https://cd.epic.epd.gov.hk/EPICDI/air/station/?lang=en",
        "8. Census and Statistics Department. Population by Sex and Age Group (Table 110-01001). https://data.gov.hk/en-data/dataset/hk-censtatd-tablechart-110-01001",
        "9. Ye X, Wolff R, Yu W, Vaneckova P, Pan X, Tong S. Ambient temperature and morbidity: a review of epidemiological evidence. Environ Health Perspect. 2012;120(1):19-28. doi:10.1289/ehp.1003198",
        "10. Bhaskaran K, Gasparrini A, Hajat S, Smeeth L, Armstrong B. Time series regression studies in environmental epidemiology. Int J Epidemiol. 2013;42(4):1187-1195. doi:10.1093/ije/dyt092",
        "11. Goggins WB, Woo J, Ho S, Chan EYY, Chau PH. Weather, season, and daily stroke admissions in Hong Kong. Int J Biometeorol. 2012;56(5):865-872. doi:10.1007/s00484-011-0491-9",
        "12. Liu J, Hansen A, Varghese B, et al. Cause-specific mortality attributable to cold and hot ambient temperatures in Hong Kong: a time-series study, 2006–2016. Sustain Cities Soc. 2020;57:102131. doi:10.1016/j.scs.2020.102131",
        "13. Liu Z, Ren C, Liu J, Kawasaki Y, Bishai DM. Modelling the excess mortality associated with heat waves in Hong Kong: 2014–2023. medRxiv. 2026. doi:10.64898/2026.03.05.26347683",
        "14. Greenland S, Morgenstern H. Ecological bias, confounding, and effect modification. Int J Epidemiol. 1989;18(1):269-274. doi:10.1093/ije/18.1.269",
        "15. Basagaña X, Ballester J. Unbiased temperature-related mortality estimates using weekly and monthly health data: a new method for environmental epidemiology and climate impact studies. Lancet Planet Health. 2024;8(10):e766-e777. doi:10.1016/S2542-5196(24)00212-2",
        "16. Basagaña X, Ballester J. Unbiased estimates using temporally aggregated outcome data in time series analysis: generalization to different outcomes, exposures, and types of aggregation. Epidemiology. 2026;37(1):16-20. doi:10.1097/EDE.0000000000001923",
        "17. Guo YT, Chan KH, Qiu H, Wong ELY, Ho KF. The risk of hospitalization associated with hot nights and excess nighttime heat in a subtropical metropolis: a time-series study in Hong Kong, 2000–2019. Lancet Reg Health West Pac. 2024;51:101168. doi:10.1016/j.lanwpc.2024.101168",
        "18. Yang Z, Wei Y, Jiang X, et al. Association of cold weather and influenza infection with stroke: a 22-year time-series analysis. Int J Biometeorol. 2025;69(5):963-973. doi:10.1007/s00484-025-02870-2",
        "19. Newey WK, West KD. A simple, positive semi-definite, heteroskedasticity and autocorrelation consistent covariance matrix. Econometrica. 1987;55(3):703-708. doi:10.2307/1913610",
        "20. Benjamini Y, Hochberg Y. Controlling the false discovery rate: a practical and powerful approach to multiple testing. J R Stat Soc Series B. 1995;57(1):289-300. doi:10.1111/j.2517-6161.1995.tb02031.x",
        "21. Goggins WB, Chan EYY. A study of the short-term associations between hospital admissions and mortality from heart failure and meteorological variables in Hong Kong. Int J Cardiol. 2017;228:537-542. doi:10.1016/j.ijcard.2016.11.106",
    ]
    for r in refs:
        p = doc.add_paragraph()
        set_paragraph_format(p, first_line=False, space_after=4)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        run = p.add_run(r)
        set_run_font(run, size=11)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")
    # Hogan weather check
    from docx import Document as D
    d = D(str(OUT))
    texts = "\n".join(p.text for p in d.paragraphs)
    assert "Meteorological data was obtained from the HKO." in texts
    assert "Acute myocardial infarction and stroke may respond differently" not in texts
    assert "housing" not in texts.lower()
    assert "medication" not in texts.lower() or "medication" in texts  # acknowledgements shouldn't have it
    assert texts.lower().count("medication") == 0
    assert "Yang CY" in texts
    print("checks passed")


if __name__ == "__main__":
    build()
