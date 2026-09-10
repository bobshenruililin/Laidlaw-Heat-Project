#!/usr/bin/env python3
"""Build the COVID-period candidate manuscript (Word, then PDF if LibreOffice).

Wording authority: manuscript/covid_period/Manuscript_covid_period_draft.md
Hogan's weather paragraph is copied verbatim, including the averaging sentence.
Does not write into manuscript/live_collaborative/.
Does not invent HA coefficients.
"""
from __future__ import annotations

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT_DOCX = ROOT / "manuscript/covid_period/Heat_CVD_Manuscript_covid_period.docx"
OUT_PDF = ROOT / "manuscript/covid_period/Heat_CVD_Manuscript_covid_period.pdf"
MS = ROOT / "manuscript/covid_period/Manuscript_covid_period_draft.md"
FIGS = {
    1: ROOT / "figures/live_identification/figure_B_first_event_depletion.png",
    2: ROOT / "figures/live_identification/figure_A_cold_day_identification.png",
    3: ROOT / "figures/live_identification/figure_D_trend_depletion_sensitivity.png",
}

HOGAN_OPEN = "Meteorological data was obtained from the HKO."
HOGAN_AVG = (
    "All monthly data were derived by taking the average of daily data in each calendar month."
)

TOKEN = re.compile(
    r"(\*\*.+?\*\*|T~min~|T~max~|NO~2~|SO~2~|O~3~|PM~2\.5~|\^[0-9]+\^|"
    r"\\\(.+?\\\)|`[^`]+`|\*[qp]\*)"
)


def set_run_font(
    run,
    *,
    italic=False,
    bold=False,
    size=12,
    color=None,
    superscript=False,
    subscript=False,
):
    run.italic = italic
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    run.font.superscript = superscript
    run.font.subscript = subscript


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    set_run_font(run, size=10, italic=True)


def set_paragraph_format(
    p, *, first_line=True, space_after=8, space_before=0, align="left"
):
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = Cm(1.27) if first_line else Cm(0)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def heading(doc, text, level=1, *, new_page=False):
    p = doc.add_paragraph()
    set_paragraph_format(
        p,
        first_line=False,
        space_before=0 if new_page else (12 if level == 1 else 8),
        space_after=6,
    )
    if new_page:
        p.paragraph_format.page_break_before = True
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, bold=True, size=14)
    elif level == 2:
        set_run_font(run, bold=True, size=12)
    else:
        set_run_font(run, bold=True, italic=True, size=12)
    return p


def _set_cell_margins(cell, top=40, bottom=40, left=60, right=60):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def _row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _row_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


def add_table(doc, headers, rows, font_size=9):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for child in list(tbl_pr):
        if child.tag == qn("w:tblW"):
            tbl_pr.remove(child)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), "5000")
    tbl_w.set(qn("w:type"), "pct")
    tbl_pr.append(tbl_w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    n_rows = 1 + len(rows)
    for r_i, row in enumerate(table.rows):
        _row_cant_split(row)
        if r_i == 0:
            _row_repeat_header(row)
        last = r_i == n_rows - 1
        vals = headers if r_i == 0 else rows[r_i - 1]
        for c_i, val in enumerate(vals):
            cell = row.cells[c_i]
            cell.text = ""
            _set_cell_margins(cell)
            p = cell.paragraphs[0]
            set_paragraph_format(p, first_line=False, space_after=1, space_before=1)
            p.paragraph_format.keep_together = True
            if not last:
                p.paragraph_format.keep_with_next = True
            run = p.add_run(str(val))
            set_run_font(run, bold=(r_i == 0), size=font_size)
    spacer = doc.add_paragraph()
    set_paragraph_format(spacer, first_line=False, space_after=6, space_before=0)


def add_picture(doc, path, width_in=6.2, *, max_height_in=7.6):
    from PIL import Image

    with Image.open(path) as im:
        px_w, px_h = im.size
    height_in = width_in * (px_h / px_w)
    if height_in > max_height_in:
        width_in = max_height_in * (px_w / px_h)
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, space_after=4, align="center")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    return p


def add_rich_runs(p, text, *, size=12, italic=False):
    subs = {
        "T~min~": ("T", "min"),
        "T~max~": ("T", "max"),
        "NO~2~": ("NO", "2"),
        "SO~2~": ("SO", "2"),
        "O~3~": ("O", "3"),
        "PM~2.5~": ("PM", "2.5"),
    }
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            run = p.add_run(text[pos : m.start()])
            set_run_font(run, size=size, italic=italic)
        tok = m.group(0)
        if tok in subs:
            base, sub = subs[tok]
            r1 = p.add_run(base)
            set_run_font(r1, size=size, italic=italic)
            r2 = p.add_run(sub)
            set_run_font(r2, size=size, italic=italic, subscript=True)
        elif tok.startswith("**") and tok.endswith("**"):
            run = p.add_run(tok[2:-2])
            set_run_font(run, size=size, bold=True, italic=italic)
        elif tok.startswith("^") and tok.endswith("^"):
            run = p.add_run(tok[1:-1])
            set_run_font(run, size=size, superscript=True)
        elif tok.startswith(r"\(") and tok.endswith(r"\)"):
            inner = tok[2:-2]
            inner = re.sub(r"\\mathrm\{([^}]+)\}", r"\1", inner)
            inner = inner.replace("\\,", "")
            run = p.add_run(inner)
            set_run_font(run, size=size, italic=True)
        elif tok in ("*q*", "*p*"):
            run = p.add_run(tok[1])
            set_run_font(run, size=size, italic=True)
        elif tok.startswith("`") and tok.endswith("`"):
            run = p.add_run(tok[1:-1])
            set_run_font(run, italic=True, size=size)
        else:
            run = p.add_run(tok)
            set_run_font(run, size=size, italic=italic)
        pos = m.end()
    if pos < len(text):
        run = p.add_run(text[pos:])
        set_run_font(run, size=size, italic=italic)


def add_body(doc, text, *, first_line=True, italic=False, size=12, align="justify"):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=first_line, space_after=8, align=align)
    add_rich_runs(p, text, size=size, italic=italic)
    return p


def _strip_md_cell(cell: str) -> str:
    return cell.replace("**", "").replace("*", "").strip()


def _parse_table(lines: list[str]) -> tuple[list[str], list[list[str]]]:
    body_lines = [ln for ln in lines if set(ln.strip()) - set("|:- ")]
    rows = []
    for ln in body_lines:
        parts = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append([_strip_md_cell(c) for c in parts])
    return rows[0], rows[1:]


def _ensure_figures() -> None:
    missing = [p for p in FIGS.values() if not p.is_file()]
    if missing:
        subprocess.run(
            ["python3", str(ROOT / "scripts" / "76_covid_period_identification_figures.py")],
            cwd=ROOT,
            check=True,
        )
    for path in FIGS.values():
        if not path.is_file():
            raise SystemExit(f"missing figure {path}")


def _blocks(md: str) -> list[tuple[str, list[str]]]:
    lines = md.splitlines()
    blocks: list[tuple[str, list[str]]] = []
    i = 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("# "):
            blocks.append(("h0", [ln[2:].strip()]))
            i += 1
            continue
        if ln.startswith("## "):
            blocks.append(("h1", [ln[3:].strip()]))
            i += 1
            continue
        if ln.startswith("### "):
            blocks.append(("h2", [ln[4:].strip()]))
            i += 1
            continue
        if ln.startswith("!["):
            blocks.append(("img", [ln]))
            i += 1
            continue
        if ln.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].startswith("|"):
                tbl.append(lines[i])
                i += 1
            blocks.append(("table", tbl))
            continue
        if ln.startswith("\\["):
            eq = [ln]
            i += 1
            while i < len(lines) and not lines[i].strip().endswith("\\]"):
                eq.append(lines[i])
                i += 1
            if i < len(lines):
                eq.append(lines[i])
                i += 1
            blocks.append(("eq", eq))
            continue
        para = [ln]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(
            ("#", "|", "![", "\\[")
        ):
            para.append(lines[i])
            i += 1
        blocks.append(("p", para))
    return blocks


def build_docx() -> None:
    md = MS.read_text(encoding="utf-8")
    assert HOGAN_OPEN in md
    assert HOGAN_AVG in md
    assert "Gate 3" not in md
    assert "placeholder" not in md.lower()
    assert "Hogan asked" not in md
    _ensure_figures()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.clear()
    set_paragraph_format(hp, first_line=False, space_after=0)
    r = hp.add_run("Window dependence of first CHD/HF hospitalisation")
    set_run_font(r, italic=True, size=10, color=RGBColor(0x55, 0x55, 0x55))
    r2 = hp.add_run("\t")
    set_run_font(r2, size=10)
    add_page_number(hp)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), "9360")
    tabs.append(tab)
    hp._p.get_or_add_pPr().append(tabs)

    saw_intro = False
    for kind, lines in _blocks(md):
        if kind == "h0":
            t = doc.add_paragraph()
            set_paragraph_format(t, first_line=False, space_after=6, align="center")
            tr = t.add_run(lines[0])
            set_run_font(tr, bold=True, size=16)
            continue
        if kind == "h1":
            title = lines[0]
            new_page = title == "Introduction" and not saw_intro
            if title == "Introduction":
                saw_intro = True
            heading(doc, title, 1, new_page=new_page)
            continue
        if kind == "h2":
            heading(doc, lines[0], 2)
            continue
        if kind == "table":
            headers, rows = _parse_table(lines)
            add_table(doc, headers, rows)
            continue
        if kind == "eq":
            text = " ".join(ln.strip() for ln in lines)
            text = text.replace("\\[", "").replace("\\]", "").replace("\\tag{1}", "(1)")
            text = re.sub(r"\\mathrm\{([^}]+)\}", r"\1", text)
            text = text.replace("\\log", "log ").replace("\\mathrm E", "E")
            text = re.sub(r"\\sum_\{[^}]+\}", "sum ", text)
            text = text.replace("\\,", " ").replace("~", " ")
            p = doc.add_paragraph()
            set_paragraph_format(p, first_line=False, space_after=8, align="center")
            run = p.add_run(text)
            set_run_font(run, italic=True, size=11)
            continue
        if kind == "img":
            m = re.search(r"!\[([^\]]*)\]\(([^)]+)\)", lines[0])
            if not m:
                continue
            alt = m.group(1)
            n = 1
            if "Figure 2" in alt:
                n = 2
            elif "Figure 3" in alt:
                n = 3
            add_picture(doc, FIGS[n])
            continue
        text = " ".join(ln.strip() for ln in lines)
        if text.startswith("*") and text.endswith("*") and not text.startswith("**"):
            add_body(doc, text.strip("*"), first_line=False, italic=True, align="center")
            continue
        if text.startswith("**Figure") or text.startswith("**Table"):
            add_body(doc, text.replace("**", ""), first_line=False, italic=True, size=11)
            continue
        first = not text.startswith("**")
        add_body(doc, text, first_line=first)

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(f"wrote {OUT_DOCX} ({OUT_DOCX.stat().st_size} bytes)")


def convert_pdf() -> None:
    cmd = [
        "soffice",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(OUT_DOCX.parent),
        str(OUT_DOCX),
    ]
    try:
        subprocess.run(cmd, check=True, capture_output=True, text=True)
    except (FileNotFoundError, subprocess.CalledProcessError) as exc:
        print(f"LibreOffice PDF skipped: {exc}")
        return
    if OUT_PDF.is_file():
        print(f"wrote {OUT_PDF} ({OUT_PDF.stat().st_size} bytes)")


def main() -> None:
    build_docx()
    convert_pdf()


if __name__ == "__main__":
    main()
