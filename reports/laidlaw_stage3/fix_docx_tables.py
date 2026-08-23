#!/usr/bin/env python3
"""Force Stage 3 tables to page width with manuscript-style rules (no wrap/overflow)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips, Cm, RGBColor


CONTENT_DXA = 9072  # ~16 cm at 1440 twips/inch


def _set(el, **attrs):
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    return el


def set_cell_border(cell, **edges):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge, spec in edges.items():
        el = tcBorders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcBorders.append(el)
        for k, v in spec.items():
            el.set(qn(f"w:{k}"), str(v))


def set_tbl_width(table, dxa):
    tbl = table._tbl
    tblPr = tbl.tblPr
    # drop broken style name "Table"
    for child in list(tblPr):
        if child.tag == qn("w:tblStyle"):
            tblPr.remove(child)
        if child.tag == qn("w:tblW"):
            tblPr.remove(child)
        if child.tag == qn("w:jc"):
            tblPr.remove(child)
    tblW = OxmlElement("w:tblW")
    _set(tblW, type="dxa", w=str(dxa))
    tblPr.append(tblW)
    jc = OxmlElement("w:jc")
    _set(jc, val="center")
    tblPr.append(jc)
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    _set(layout, type="fixed")


def set_grid(table, widths):
    tbl = table._tbl
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        col = OxmlElement("w:gridCol")
        _set(col, w=str(w))
        grid.append(col)
    # insert after tblPr
    tblPr = tbl.tblPr
    tblPr.addnext(grid)


def set_cell_width(cell, dxa):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    _set(tcW, type="dxa", w=str(dxa))
    mar = tcPr.find(qn("w:tcMar"))
    if mar is not None:
        tcPr.remove(mar)
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", "40"), ("left", "60"), ("bottom", "40"), ("right", "60")):
        el = OxmlElement(f"w:{edge}")
        _set(el, w=val, type="dxa")
        mar.append(el)
    tcPr.append(mar)


def style_run(paragraph, size=10, bold=False, align=None):
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for run in paragraph.runs:
        run.font.name = "Liberation Serif"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(0, 0, 0)
        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), "Liberation Serif")
        rFonts.set(qn("w:hAnsi"), "Liberation Serif")


THIN = {"val": "single", "sz": "4", "space": "0", "color": "000000"}
NONE = {"val": "nil"}
THICK = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
MID = {"val": "single", "sz": "8", "space": "0", "color": "000000"}


def proportion_widths(n, fractions):
    assert len(fractions) == n
    raw = [int(CONTENT_DXA * f) for f in fractions]
    raw[-1] += CONTENT_DXA - sum(raw)
    return raw


TABLE_FRACTIONS = {
    4: (0.34, 0.14, 0.26, 0.26),
    5: (0.28, 0.18, 0.18, 0.18, 0.18),
}


def fix_table(table):
    n = len(table.columns)
    fracs = TABLE_FRACTIONS.get(n)
    if fracs is None:
        fracs = tuple([1 / n] * n)
    # special-case 4-col core panel (outcome short, exposure long)
    first = table.rows[0].cells[0].text.strip()
    second = table.rows[0].cells[1].text.strip() if n > 1 else ""
    if n == 4 and "Uncertainty" in second:
        fracs = (0.36, 0.26, 0.14, 0.24)
    if n == 4 and "Exposure" in second:
        fracs = (0.12, 0.42, 0.32, 0.14)
    if n == 4 and first == "Outcome" and "Months" in table.rows[0].cells[1].text:
        fracs = (0.40, 0.16, 0.22, 0.22)
    widths = proportion_widths(n, fracs)
    set_tbl_width(table, CONTENT_DXA)
    set_grid(table, widths)
    nrows = len(table.rows)
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            set_cell_width(cell, widths[ci])
            if ri == 0:
                borders = dict(top=THICK, bottom=MID, left=NONE, right=NONE)
            elif ri == nrows - 1:
                borders = dict(top=NONE, bottom=THICK, left=NONE, right=NONE)
            else:
                borders = dict(top=NONE, bottom=NONE, left=NONE, right=NONE)
            set_cell_border(cell, **borders)
            align = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 or (n >= 4 and ci == 1 and "Exposure" in second) else WD_ALIGN_PARAGRAPH.RIGHT
            if ri == 0:
                align = WD_ALIGN_PARAGRAPH.LEFT if ci <= 1 else WD_ALIGN_PARAGRAPH.RIGHT
            for p in cell.paragraphs:
                style_run(p, size=10, bold=(ri == 0), align=align)


def main():
    path = Path("reports/laidlaw_stage3/ShenRuililin_Laidlaw_Stage3Report.docx")
    doc = Document(str(path))
    for t in doc.tables:
        fix_table(t)
    # captions: keep with next, no wrap
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("Table ") or t.startswith("Figure "):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(10)
                run.font.name = "Liberation Serif"
        if t.startswith("Note.") or t.startswith("*Note."):
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(8)
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.italic = True
    doc.save(str(path))
    print("fixed tables in", path)


if __name__ == "__main__":
    main()
